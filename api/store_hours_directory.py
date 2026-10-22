"""Opening/closing store directory and validated two-sheet schedule uploads."""
from datetime import datetime, time
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
from zipfile import ZipFile, BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from api.calling_tree import active_roster
from api.portal import PortalError
from api.store_schedules import scheduled_hours

DEALER_ALIASES = {'CONNECT': 'Connect', 'CONNECT WIRELESS': 'Connect',
                  'SUPREME': 'SRH', 'SRH': 'SRH', 'ARM WIRELESS': 'AMQ',
                  'DF WIRELESS': 'AMQ', 'AMQ': 'AMQ', 'TALKALOT WIRELESS': 'ARM',
                  'ARM': 'ARM', 'ARBF': 'ARBF', 'ARBF WIRELESS METRO': 'ARBF',
                  'THE FIX': 'The Fix', 'SMOKE MAGIC': 'Smoke Magic',
                  'SPDI-CA': 'California', 'CALIFORNIA': 'California'}
SOURCE_IDS = {'Connect': 'connect', 'California': 'california', 'SRH': 'srh',
              'AMQ': 'arm1', 'ARM': 'arm2', 'ARBF': 'arbf'}
DEALER_LABELS = {'AMQ': 'AMQ / ARM Wireless', 'ARM': 'ARM / Talkalot Wireless',
                 'SRH': 'SRH / Supreme', 'Connect': 'Connect Wireless'}
MAX_FILE_BYTES = 2 * 1024 * 1024


def normalized(value):
    return ' '.join(str(value or '').split()).upper()


def roster_with_schedules(conn):
    stores = {}
    for row in active_roster(conn):
        stores[row['dealer'], row['store_id']] = {**row, 'source_id': SOURCE_IDS.get(row['dealer'])}
    for row in conn.execute('select * from public.store_hours_directory order by dealer,market,store').fetchall():
        key = row['dealer'], row['store_id']
        # Calling Tree names stay current; the workbook adds schedules and source mapping.
        stores[key] = {**row, **stores.get(key, {}), 'schedules': row['schedules'],
                       'source_id': row['source_id'], 'company': row['company']}
    return sorted(stores.values(), key=lambda s: (s['dealer'], s['market'], s['store']))


def hours_for_store(store, day):
    if 'schedules' not in store:
        return scheduled_hours(store, day)
    period = 'sunday' if day.weekday() == 6 else 'saturday' if day.weekday() == 5 else 'weekday'
    label = {'weekday': 'Monday–Friday', 'saturday': 'Saturday', 'sunday': 'Sunday'}[period]
    return {**{k: time.fromisoformat(v) for k, v in store['schedules'].get(period, {}).items()}, 'period': label}


def parse_time(value, location):
    if isinstance(value, datetime):
        value = value.time()
    if isinstance(value, time) and value.tzinfo is None:
        return value.replace(microsecond=0).isoformat()
    if isinstance(value, str):
        for fmt in ('%I:%M:%S %p', '%I:%M %p', '%H:%M:%S', '%H:%M'):
            try:
                return datetime.strptime(value.strip().upper(), fmt).time().isoformat()
            except ValueError:
                pass
    raise PortalError(f'{location}: enter a valid opening/closing time (for example 10:00 AM).')


def parse_schedule_workbook(data):
    if len(data) > MAX_FILE_BYTES:
        raise PortalError('Choose an .xlsx workbook smaller than 2 MB.')
    try:
        with ZipFile(BytesIO(data)) as archive:
            if len(archive.infolist()) > 2000 or sum(i.file_size for i in archive.infolist()) > 25 * 1024 * 1024:
                raise PortalError('This workbook is too large when unpacked.')
            if any('vbaproject' in i.filename.lower() for i in archive.infolist()):
                raise PortalError('Upload an .xlsx workbook without macros.')
        book = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
    except (BadZipFile, InvalidFileException, OSError, ValueError, KeyError):
        raise PortalError('Choose a valid .xlsx schedule workbook.') from None
    try:
        sheets = {normalized(s.title): s for s in book}
        if not {'MON TO SAT', 'SUN'} <= sheets.keys():
            raise PortalError('The workbook needs both MON TO SAT and SUN sheets.')
        result = {}
        for name, width in (('MON TO SAT', 8), ('SUN', 6)):
            sheet = sheets[name]
            if (sheet.max_row or 0) > 5002 or (sheet.max_column or 0) > 20:
                raise PortalError('Use at most 5,000 stores and 20 columns per sheet.')
            expected = ['CARRIER', 'COMPANY NAME', 'MARKET', 'STORE', 'OPENING TIME', 'CLOSING TIME']
            if name == 'MON TO SAT':
                expected += ['OPENING TIME', 'CLOSING TIME']
            headers = next(sheet.iter_rows(min_row=2, max_row=2, max_col=width, values_only=True))
            if [normalized(v) for v in headers] != expected:
                raise PortalError(f'{name}: preserve the original headings in row 2.')
            seen = set()
            for n, cells in enumerate(sheet.iter_rows(min_row=3, max_col=width), 3):
                values = [c.value for c in cells]
                if not any(v is not None and str(v).strip() for v in values):
                    continue
                if any(c.data_type in ('f', 'e') for c in cells):
                    raise PortalError(f'{name} row {n}: replace formulas and errors with plain values.')
                identity = tuple(normalized(v) for v in values[:4])
                if not all(identity) or any(len(v) > 160 for v in identity):
                    raise PortalError(f'{name} row {n}: Carrier, Company Name, Market and Store are required (up to 160 characters).')
                if identity in seen:
                    raise PortalError(f'{name} row {n}: duplicate store.')
                seen.add(identity)
                entry = result.setdefault(identity, dict(zip(('carrier', 'company', 'market', 'store'), identity)))
                for period, offset in ([('weekday', 4), ('saturday', 6)] if name == 'MON TO SAT' else [('sunday', 4)]):
                    opening = parse_time(values[offset], f'{name} row {n}')
                    closing = parse_time(values[offset+1], f'{name} row {n}')
                    # Preserve the user's explicit correction for this workbook format.
                    if DEALER_ALIASES.get(identity[1]) == 'Connect' and period != 'sunday' and opening == '22:00:00':
                        opening = '10:00:00'
                        entry['corrected_connect_opening'] = True
                    entry.setdefault('schedules', {})[period] = {'opening_time': opening, 'closing_time': closing}
        if not result:
            raise PortalError('The workbook contains no stores.')
        for entry in result.values():
            if set(entry['schedules']) != {'weekday', 'saturday', 'sunday'}:
                raise PortalError(f"{entry['store']}: list this store with the same carrier, company and market on both sheets.")
        return list(result.values())
    finally:
        book.close()


def map_schedule_rows(rows, roster):
    seed = json.loads(Path(__file__).with_name('store_schedules.json').read_text(encoding='utf-8'))['stores']
    aliases = {(normalized(s['company']), normalized(s['market']), normalized(s['store'])):
               (s['portal_dealer'], s['portal_store_id']) for s in seed if s['portal_store_id']}
    mapped, keys = [], set()
    for row in rows:
        dealer = DEALER_ALIASES.get(row['company'], row['company'].title())
        preferred = aliases.get((row['company'], row['market'], row['store']))
        matches = [s for s in roster if
                   (preferred and (s['dealer'], s['store_id']) == preferred) or
                   (s['dealer'] == dealer and normalized(s['market']) == row['market'] and normalized(s['store']) == row['store'])]
        if len(matches) > 1:
            raise PortalError(f"{row['store']}: ambiguous store match. Use the exact company, market and store names.")
        if matches:
            existing = matches[0]
            dealer, sid = existing['dealer'], existing['store_id']
            source = existing.get('source_id')
        else:
            # Include carrier to distinguish ARBF Total from ARBF Metro.
            identity = [dealer, row['carrier'], row['market'], row['store']]
            sid = 'WB-' + sha256(json.dumps(identity).encode()).hexdigest()[:24]
            source = None
        if (dealer, sid) in keys:
            raise PortalError(f"{row['store']}: multiple workbook rows map to the same store.")
        keys.add((dealer, sid))
        mapped.append({**row, 'dealer': dealer, 'store_id': sid, 'source_id': source,
                       'action': 'Update' if matches else 'Add'})
    return mapped
