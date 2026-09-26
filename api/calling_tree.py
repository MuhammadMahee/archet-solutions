"""Validated, versioned Calling Tree uploads. Workbook cells are data, never commands."""
import base64
import binascii
from collections import Counter
from io import BytesIO
from pathlib import PureWindowsPath
import re
from uuid import UUID
from zipfile import ZipFile, BadZipFile

from flask import request, jsonify, g
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

try:
    from api.portal import require_user, PortalError, body
except ModuleNotFoundError:
    from portal import require_user, PortalError, body

MAX_FILE_BYTES = 2 * 1024 * 1024
FIELDS = {'Dealer':'original_dealer','Store ID':'store_id','Market':'market','Store Name':'store',
          'Carrier':'carrier','DM':'dm','State':'state','Dealer Code':'dealer_code',
          'Door Code':'door_code','SAP ID':'sap_id','Address':'address','ZIP Code':'zip_code'}
REQUIRED = {'Dealer','Store ID','Market','Store Name'}
ALIASES = {'connect':'Connect','california':'California','spdica':'California',
           'srh':'SRH','supreme':'SRH','amq':'AMQ','arm':'ARM',
           'arm1':'AMQ','arm2':'ARM','armwireless':'ARM','dfwireless':'AMQ',
           'arbf':'ARBF','arbfwirelessmetro':'ARBF'}
DEALERS = {'connect':'Connect','california':'California','srh':'SRH','arm1':'AMQ','arm2':'ARM','arbf':'ARBF'}


def clean(value):
    if value is None:
        return ''
    if isinstance(value,float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def parse_workbook(data, worksheet=''):
    if len(data)>MAX_FILE_BYTES:
        raise PortalError('Choose an .xlsx workbook smaller than 2 MB.')
    try:
        with ZipFile(BytesIO(data)) as archive:
            if len(archive.infolist())>2000 or sum(i.file_size for i in archive.infolist())>25*1024*1024:
                raise PortalError('This workbook is too large when unpacked.')
            if any('vbaproject' in i.filename.lower() for i in archive.infolist()):
                raise PortalError('Upload an .xlsx workbook without macros.')
        book=load_workbook(BytesIO(data),read_only=True,data_only=False,keep_links=False)
    except (BadZipFile, InvalidFileException, OSError, ValueError, KeyError):
        raise PortalError('The file could not be read. Upload a valid .xlsx workbook.') from None
    try:
        candidates=[]
        if worksheet and worksheet not in book.sheetnames:
            raise PortalError('That worksheet was not found in the workbook.')
        for sheet in book:
            if worksheet and sheet.title!=worksheet:
                continue
            if (sheet.max_column or 0)>100 or (sheet.max_row or 0)>5021:
                raise PortalError('Each worksheet must contain at most 5,000 stores and 100 columns.')
            for idx,row in enumerate(sheet.iter_rows(min_row=1,max_row=20),1):
                headers={clean(c.value).casefold():i for i,c in enumerate(row) if c.value is not None}
                if {h.casefold() for h in REQUIRED}<=headers.keys():
                    candidates.append((sheet,idx,headers));break
        if not candidates:
            raise PortalError('Required columns: Dealer, Store ID, Market, Store Name.')
        if len(candidates)>1:
            raise PortalError('Several worksheets contain store lists. Enter the worksheet name to import.')
        sheet,header_row,headers=candidates[0]
        result=[];seen=set()
        for idx,cells in enumerate(sheet.iter_rows(min_row=header_row+1),header_row+1):
            mapped={}
            for heading,key in FIELDS.items():
                col=headers.get(heading.casefold())
                cell=cells[col] if col is not None and col<len(cells) else None
                if cell and cell.data_type in ('f','e'):
                    raise PortalError(f'Row {idx}: replace formulas or spreadsheet errors with plain values.')
                value=clean(cell.value if cell else None)
                if len(value)>500:
                    raise PortalError(f'Row {idx}: {heading} is too long.')
                if heading=='ZIP Code' and cell and re.fullmatch(r'0{5}',cell.number_format or '') and value.isdigit():
                    value=value.zfill(5)
                mapped[key]=value
            if not any(mapped.values()):
                continue
            if any(not mapped[FIELDS[key]] for key in REQUIRED):
                raise PortalError(f'Row {idx}: Dealer, Store ID, Market and Store Name cannot be empty.')
            alias=re.sub(r'[^a-z0-9]','',mapped['original_dealer'].lower())
            dealer=ALIASES.get(alias)
            if not dealer:
                raise PortalError(f'Row {idx}: unknown dealer. Use Connect, California, SRH, AMQ, ARM or ARBF.')
            mapped.update(dealer=dealer,store_id=mapped['store_id'].upper(),
                          market=mapped['market'].upper(),store=mapped['store'].upper(),row_number=idx)
            key=(dealer,mapped['store_id'])
            if key in seen:
                raise PortalError(f'Row {idx}: duplicate Store ID within the same dealer.')
            seen.add(key);result.append(mapped)
            if len(result)>5000:
                raise PortalError('Upload at most 5,000 stores.')
        if not result:
            raise PortalError('The worksheet contains no stores.')
        return sheet.title,result
    finally:
        book.close()


def active_roster(conn):
    # Single statement gives one coherent version and store list during replacement.
    records=conn.execute('''select s.*,v.filename,v.worksheet,v.activated_at from public.calling_tree_stores s
        join public.calling_tree_versions v on v.id=s.version_id where v.status='active'
        order by s.row_number''').fetchall()
    return records


def source_catalog(conn):
    catalog={}
    for r in conn.execute('select distinct source_id,store_id from public.sales_performance').fetchall():
        catalog.setdefault((DEALERS[r['source_id']],r['store_id']),[]).append(r['source_id'])
    return catalog


def describe_rows(rows,catalog):
    return [{**r,'matched':(r['dealer'],r['store_id']) in catalog} for r in rows]


def register_calling_tree(app,connect):
    @app.get('/api/internal/calling-tree')
    @require_user()
    def get_tree():
        with connect() as conn:
            rows=active_roster(conn)
            catalog=source_catalog(conn)
        first=rows[0] if rows else None
        return jsonify(rows=describe_rows(rows,catalog),
            active={'id':first['version_id'],'filename':first['filename'],'worksheet':first['worksheet'],
                    'activated_at':first['activated_at'],'row_count':len(rows)} if first else None)

    @app.post('/api/internal/calling-tree/preview')
    @require_user(admin=True)
    def preview_tree():
        payload=body()
        encoded=payload.get('content','');filename=payload.get('filename','');worksheet=payload.get('worksheet','')
        if not all(isinstance(v,str) for v in (encoded,filename,worksheet)) or len(encoded)>2800000 or len(filename)>250 or len(worksheet)>100:
            raise PortalError('Choose a valid workbook smaller than 2 MB.')
        filename=PureWindowsPath(filename).name
        if not filename.lower().endswith('.xlsx'):
            raise PortalError('Upload an .xlsx workbook.')
        try:
            data=base64.b64decode(encoded,validate=True)
        except (ValueError,binascii.Error):
            raise PortalError('The workbook upload was incomplete. Choose the file again.') from None
        worksheet,rows=parse_workbook(data,worksheet)
        with connect() as conn:
            with conn.transaction():
                conn.execute('select pg_advisory_xact_lock(731044)')
                # Expired previews contain no active roster data and need not accumulate.
                conn.execute("delete from public.calling_tree_versions where status='draft' and uploaded_at<now()-interval '1 day'")
                active=conn.execute("select id from public.calling_tree_versions where status='active'").fetchone()
                result=conn.execute('''insert into public.calling_tree_versions(filename,worksheet,uploaded_by,base_version_id,row_count)
                    values(%s,%s,%s,%s,%s) returning id''',
                    (filename,worksheet,g.portal_user['id'],active['id'] if active else None,len(rows))).fetchone()
                version=result['id'];fields=('row_number',*FIELDS.values(),'dealer')
                with conn.cursor() as cur:
                    cur.executemany(f'''insert into public.calling_tree_stores(version_id,{','.join(fields)})
                        values({','.join(['%s']*(len(fields)+1))})''',[(version,*(r[k] for k in fields)) for r in rows])
                catalog=source_catalog(conn)
        described=describe_rows(rows,catalog)
        return jsonify(preview_id=version,filename=filename,worksheet=worksheet,rows=described,
                       counts=dict(Counter(r['dealer'] for r in rows)),unmatched=sum(not r['matched'] for r in described))

    @app.post('/api/internal/calling-tree/activate')
    @require_user(admin=True)
    def activate_tree():
        try:
            version=str(UUID(str(body().get('preview_id',''))))
        except ValueError:
            raise PortalError('Preview the workbook before applying it.') from None
        with connect() as conn:
            with conn.transaction():
                conn.execute('select pg_advisory_xact_lock(731044)')
                draft=conn.execute("select * from public.calling_tree_versions where id=%s and status='draft' and uploaded_by=%s",
                    (version,g.portal_user['id'])).fetchone()
                if not draft:
                    raise PortalError('This preview is unavailable. Upload the workbook again.',409)
                current=conn.execute("select id from public.calling_tree_versions where status='active'").fetchone()
                if draft['base_version_id']!=(current['id'] if current else None):
                    raise PortalError('Another admin updated the Calling Tree. Preview your workbook again before replacing it.',409)
                conn.execute("update public.calling_tree_versions set status='archived' where status='active'")
                conn.execute("update public.calling_tree_versions set status='active',activated_at=now() where id=%s",(version,))
        return jsonify(message='Calling Tree updated. Sales Update now uses these stores.',row_count=draft['row_count'])
