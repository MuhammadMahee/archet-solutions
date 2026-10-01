"""Bounded, parallel RT-POS imports and authenticated sales reporting."""
from concurrent.futures import ThreadPoolExecutor
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from http.cookiejar import Cookie, CookieJar
from http.client import IncompleteRead
from io import BytesIO
import json
import logging
import os
import secrets
import ssl
import time
from urllib.error import HTTPError, URLError
from urllib.parse import urlencode, urlsplit
from urllib.request import Request, build_opener, HTTPCookieProcessor, HTTPRedirectHandler
from zoneinfo import ZoneInfo

from flask import jsonify, request, send_file
import psycopg
from psycopg.rows import dict_row
import xlrd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

try:
    from api.portal import require_user, PortalError
    from api.supabase_store import StoreError
    from api.calling_tree import active_roster, source_catalog
except ModuleNotFoundError:
    from portal import require_user, PortalError
    from supabase_store import StoreError
    from calling_tree import active_roster, source_catalog

REPORT_URL = 'https://myrtpos.com/newbdi/Store_Performance_lp.fwx'
CENTRAL = ZoneInfo('America/Chicago')
SOURCES = {'connect': 'Connect', 'california': 'California', 'srh': 'SRH',
           'arm1': 'AMQ', 'arm2': 'ARM', 'arbf': 'ARBF'}
COLORS = {'Connect': '#095570', 'California': '#a92d49', 'SRH': '#a06118',
          'AMQ': '#365cad', 'ARM': '#176b56', 'ARBF': '#7646a5'}
METRICS = ('new_activation', 'upgrade', 'reactivation', 'bts', 'hsi',
           'accessory', 'total_boxes', 'qpay')
COLUMNS = [('dealer', 'Dealer'), ('market', 'Market'), ('store', 'Store'),
           ('new_activation', 'New activation'), ('upgrade', 'Upgrade'),
           ('reactivation', 'Reactivation'), ('bts', 'BTS'), ('hsi', 'HSI'),
           ('accessory', 'Accessory'), ('apo', 'APO'), ('total_boxes', 'Total boxes'),
           ('qpay', 'QPay'), ('qpay_conv', 'QPay conv')]
LOCK_ID = 731043260926
MAX_REPORT_BYTES = 20 * 1024 * 1024
SOURCE_BUDGET_SECONDS = 210
SOURCE_ATTEMPTS = 3


class ReportError(Exception):
    """Only stable, non-secret error codes may leave the importer."""


def today():
    return datetime.now(CENTRAL).date()


def connect():
    url = os.getenv('SUPABASE_DB_URL', '')
    if not url:
        raise StoreError(code='not_configured')
    return psycopg.connect(url, autocommit=True, sslmode='require', connect_timeout=10,
                          row_factory=dict_row, options='-c statement_timeout=30000')


class SameHostRedirect(HTTPRedirectHandler):
    def redirect_request(self, req, fp, code, msg, headers, newurl):
        target = urlsplit(newurl)
        if target.scheme != 'https' or target.hostname not in ('myrtpos.com', 'www.myrtpos.com'):
            raise ReportError('unexpected_redirect')
        return super().redirect_request(req, fp, code, msg, headers, newurl)


def source_error(exc):
    """Return diagnostics without leaking credentials or response contents."""
    if isinstance(exc, ReportError):
        return str(exc)
    if isinstance(exc, HTTPError):
        return f'source_http_{exc.code}'
    reason = exc.reason if isinstance(exc, URLError) else exc
    if isinstance(reason, TimeoutError):
        return 'source_timeout'
    if isinstance(reason, ssl.SSLCertVerificationError):
        return 'source_tls_error'
    return 'source_unavailable'


def read_source(opener, req, limit, *, deadline=None):
    """Retry only transient transport errors within the worker's time budget."""
    deadline = deadline if deadline is not None else time.monotonic() + 110
    for attempt in range(SOURCE_ATTEMPTS):
        remaining = deadline - time.monotonic()
        if remaining <= 0:
            raise ReportError('source_timeout')
        try:
            with opener.open(req, timeout=min(35, remaining)) as response:
                return response.read(limit)
        except (OSError, IncompleteRead) as exc:
            code = source_error(exc)
            transient = (exc.code in (408, 429, 500, 502, 503, 504)
                         if isinstance(exc, HTTPError) else code != 'source_tls_error')
            if isinstance(exc, HTTPError):
                exc.close()
            logging.getLogger(__name__).warning('RT-POS %s attempt %s: %s',
                'session' if isinstance(req, str) else 'export', attempt + 1, code)
            delay = (1, 3, 0)[attempt]
            if not transient or attempt + 1 == SOURCE_ATTEMPTS or deadline - time.monotonic() <= delay:
                raise
            time.sleep(delay)


def open_source(cookies, *, deadline=None):
    """Ignore shared FW_SessionID: RT-POS must issue a NEW session per account."""
    jar = CookieJar()
    for name in ('sec85952EAF_id', 'sec85952EAF_pd'):
        value = cookies.get(name)
        if not isinstance(value, str) or not value or any(c in value for c in '\r\n;'):
            raise ReportError('missing_cookies')
        jar.set_cookie(Cookie(0, name, value, None, False, 'myrtpos.com', False,
                              False, '/newbdi', True, True, None, True, None, None, {}, False))
    opener = build_opener(SameHostRedirect(), HTTPCookieProcessor(jar))
    opener.addheaders = [('User-Agent', 'Mozilla/5.0 (compatible; ArchetSales/1.0)')]
    html = read_source(opener, REPORT_URL, 2 * 1024 * 1024, deadline=deadline).decode('utf-8', errors='replace')
    if 'frmStart' not in html or 'frmEnd' not in html:
        raise ReportError('cookies_expired')
    return opener


def fetch_report(opener, report_date, *, deadline=None):
    formatted = report_date.strftime('%m/%d/%Y')
    payload = urlencode({'frmMarketID': '', 'frmRegionID': '', 'frmStateID': '',
                         'frmStoreType': '', 'frmStore': '', 'frmStart': formatted,
                         'frmEnd': formatted, 'btnExcel': 'click'}).encode()
    req = Request(REPORT_URL, data=payload, headers={
        'Content-Type': 'application/x-www-form-urlencoded', 'Referer': REPORT_URL})
    data = read_source(opener, req, MAX_REPORT_BYTES + 1, deadline=deadline)
    if len(data) > MAX_REPORT_BYTES:
        raise ReportError('report_too_large')
    # RT-POS uses raw BIFF2 as well as OLE-wrapped XLS; xlrd validates both.
    signatures = (b'\xd0\xcf\x11\xe0', b'\x09\x00', b'\x09\x02', b'\x09\x04', b'\x09\x08')
    if not data.startswith(signatures):
        raise ReportError('invalid_report_or_expired_cookies')
    return read_xls(data)


def number(value):
    try:
        result = Decimal(str(value or '0').strip().replace(',', '').replace('$', '').replace('%', ''))
        if not result.is_finite():
            raise InvalidOperation
        return result
    except (InvalidOperation, ValueError):
        raise ReportError('invalid_numeric_data') from None


def read_xls(data):
    try:
        book = xlrd.open_workbook(file_contents=data)
        sheet = book.sheet_by_index(0)
    except (xlrd.XLRDError, IndexError):
        raise ReportError('invalid_xls') from None
    if not sheet.nrows:
        raise ReportError('missing_headers')
    headers = {str(v).strip().lower(): i for i, v in enumerate(sheet.row_values(0))}
    react = 'reactact' if 'reactact' in headers else 'react'
    needed = {'custno', 'marketid', 'company', 'newact', 'upgsor', 'totact',
              'totaccessory', 'totpaymentqty', 'iot', 'hint', react}
    if needed - headers.keys():
        raise ReportError('report_columns_changed')
    rows, seen = [], set()
    for idx in range(1, sheet.nrows):
        cell = lambda name: sheet.cell_value(idx, headers[name])
        store, sid = str(cell('company')).strip().upper(), str(cell('custno')).strip().upper()
        if not store or store == 'TOTAL':
            continue
        if not sid or sid in seen:
            raise ReportError('missing_or_duplicate_store_id')
        seen.add(sid)
        integer = lambda name: int(number(cell(name)).quantize(Decimal('1'), rounding=ROUND_HALF_UP))
        rows.append({'store_id': sid, 'store': store,
                     'market': str(cell('marketid')).strip().upper() or 'UNASSIGNED',
                     'new_activation': integer('newact'), 'upgrade': integer('upgsor'),
                     'reactivation': integer(react), 'hsi': integer('hint'),
                     'bts': integer('iot') - integer('hint'),
                     'accessory': number(cell('totaccessory')).quantize(Decimal('.01'), rounding=ROUND_HALF_UP),
                     'total_boxes': integer('totact'), 'qpay': integer('totpaymentqty')})
    # A valid header-only export means no source rows, not an authentication error.
    return rows


def seed_jobs(conn, current, *, force_today=False):
    """Resume month backfill and refresh today; close yesterday after midnight."""
    with conn.transaction():
        with conn.cursor() as cur:
            cur.executemany('''insert into public.sales_imports(source_id, report_date)
                select %s, d::date from generate_series(%s::date,%s::date,'1 day') d
                on conflict do nothing''', [(s, current.replace(day=1), current) for s in SOURCES])
        conn.execute('''update public.sales_imports set status='pending', retry_at=now()
            where status='complete' and ((report_date=%s and loaded_at < now()-interval '18 minutes')
              or (report_date=%s and loaded_at < %s))''',
                     (current, current - timedelta(days=1), datetime.combine(current, datetime.min.time(), CENTRAL)))
        if force_today:
            conn.execute('''update public.sales_imports set status='pending', retry_at=now()
                where report_date=%s or (report_date<=%s and status<>'complete')''',
                         (current, current))


def save_rows(conn, source, report_date, rows):
    """Replace explicit values, retaining omitted source rows with a stale timestamp."""
    with conn.transaction():
        fields = ('store_id', 'market', 'store', *METRICS)
        now = conn.execute('select now() as value').fetchone()['value']
        if rows:
            updates = ','.join(f'{key}=excluded.{key}' for key in fields if key != 'store_id')
            sql = f'''insert into public.sales_performance(source_id,report_date,{','.join(fields)},loaded_at)
                values ({','.join(['%s'] * (len(fields) + 3))})
                on conflict(source_id,report_date,store_id) do update set {updates},loaded_at=excluded.loaded_at'''
            with conn.cursor() as cur:
                cur.executemany(sql, [(source, report_date, *(r[k] for k in fields), now) for r in rows])
        retained = conn.execute('''select count(*) as n from public.sales_performance
            where source_id=%s and report_date=%s and loaded_at<%s''', (source, report_date, now)).fetchone()['n']
        conn.execute('''update public.sales_imports set status='complete', loaded_at=%s,
            attempted_at=%s,row_count=%s,retained_count=%s,error_code=null
            where source_id=%s and report_date=%s''', (now, now, len(rows), retained, source, report_date))


def mark_error(conn, source, report_date, code):
    conn.execute('''update public.sales_imports set status='error',attempted_at=now(),
        retry_at=now()+interval '20 minutes',error_code=%s where source_id=%s and report_date=%s''',
                 (code, source, report_date))


def source_worker(source, cookies, current):
    deadline = time.monotonic() + SOURCE_BUDGET_SECONDS
    with connect() as conn:
        jobs = conn.execute('''select report_date from public.sales_imports
            where source_id=%s and report_date<=%s and status in ('pending','error') and retry_at<=now()
            order by (report_date=%s) desc,report_date limit 4''', (source, current, current)).fetchall()
        if not jobs:
            return {'source': source, 'completed': 0, 'failed': 0}
        try:
            opener = open_source(cookies, deadline=deadline)
        except Exception as exc:
            code = source_error(exc)
            for job in jobs:
                mark_error(conn, source, job['report_date'], code)
            return {'source': source, 'completed': 0, 'failed': len(jobs), 'error': code}
        completed, failed = 0, 0
        for job in jobs:
            if time.monotonic() >= deadline:
                break  # Keep unattempted dates due for the next bounded batch.
            try:
                rows = fetch_report(opener, job['report_date'], deadline=deadline)
                save_rows(conn, source, job['report_date'], rows)
                completed += 1
            except (ReportError, OSError, IncompleteRead, ValueError) as exc:
                code = source_error(exc)
                mark_error(conn, source, job['report_date'], code)
                failed += 1
        return {'source': source, 'completed': completed, 'failed': failed}


def run_batch(*, force_today=False):
    try:
        cookies = json.loads(os.getenv('RTPOS_COOKIES', '{}'))
        if not isinstance(cookies, dict):
            raise ValueError
    except ValueError:
        raise PortalError('RT-POS credentials need configuration.', 503) from None
    with connect() as conn:
        if not conn.execute('select pg_try_advisory_lock(%s) as locked', (LOCK_ID,)).fetchone()['locked']:
            return {'status': 'busy', 'remaining': 1, 'workers': []}
        try:
            current = today()
            seed_jobs(conn, current, force_today=force_today)
            with ThreadPoolExecutor(max_workers=6) as pool:
                futures = [pool.submit(source_worker, s, cookies.get(s, {}), current) for s in SOURCES]
                results = [future.result() for future in futures]
            counts = conn.execute('''select count(*) filter(where status<>'complete') as remaining,
                count(*) filter(where status<>'complete' and retry_at<=now()) as due,
                count(*) filter(where status='error') as failed
                from public.sales_imports where report_date<=%s''', (current,)).fetchone()
            return {'status': 'partial' if counts['remaining'] else 'complete', **counts, 'workers': results}
        finally:
            conn.execute('select pg_advisory_unlock(%s)', (LOCK_ID,))


def ratios(row):
    row['apo'] = (Decimal(str(row['accessory'])) / row['total_boxes'] if row['total_boxes'] else Decimal(0)).quantize(Decimal('.01'), rounding=ROUND_HALF_UP)
    row['qpay_conv'] = (Decimal(row['total_boxes']) / row['qpay'] * 100 if row['qpay'] else Decimal(0)).quantize(Decimal('.01'), rounding=ROUND_HALF_UP)
    return row


def aggregate(raw):
    """Aggregate each dealer independently, preferring fresh copies of a store/day."""
    daily = {}
    for row in sorted(raw, key=lambda r: (r['stale'], r['source_id'])):
        key = (SOURCES[row['source_id']], row['report_date'], row['store_id'])
        daily.setdefault(key, row)
    grouped = {}
    for (dealer, day, sid), r in sorted(daily.items(), key=lambda item: item[0][1]):
        key = dealer, sid
        if key not in grouped:
            grouped[key] = {'dealer': dealer, 'store_id': sid, 'market': r['market'], 'store': r['store'],
                            'stale': False, **{k: 0 for k in METRICS}}
        out = grouped[key]
        out.update(market=r['market'], store=r['store'])
        out['stale'] |= r['stale']
        for k in METRICS:
            out[k] += r[k]
    return sorted([ratios(r) for r in grouped.values()], key=lambda r: (r['dealer'], r['market'], r['store']))


def date_range():
    try:
        start = date.fromisoformat(request.args.get('start', today().isoformat()))
        end = date.fromisoformat(request.args.get('end', start.isoformat()))
    except ValueError:
        raise PortalError('Choose valid report dates.') from None
    if start > end or end > today() or (end - start).days > 92:
        raise PortalError('Choose a range of up to 93 days ending today or earlier.')
    return start, end


def restrict_to_roster(rows, roster, catalog, imports, start, end):
    """The active Calling Tree is the allowlist, including stores with no source activity."""
    by_id={(r['dealer'],r['store_id']):r for r in rows}
    completed={(i['source_id'],i['report_date']) for i in imports if i['status']=='complete'}
    days=[start+timedelta(days=n) for n in range((end-start).days+1)]
    result=[]
    for store in roster:
        key=store['dealer'],store['store_id'];sources=catalog.get(key,[])
        available=bool(sources) and all(any((s,d) in completed for s in sources) for d in days)
        out=by_id.get(key)
        if out is None:
            out=ratios({k:0 for k in METRICS}) if available else {k:None for k in (*METRICS,'apo','qpay_conv')}
            out.update(stale=False,no_activity=available)
        else:
            out=dict(out)
        out.update(dealer=store['dealer'],store_id=store['store_id'],market=store['market'],
                   store=store['store'],incomplete=not available,unmatched=not bool(sources))
        result.append(out)
    return sorted(result,key=lambda r:(r['dealer'],r['market'],r['store']))


def metric_fills(rows):
    """Continuous red/yellow/green scales per visible column; missing data stays blank."""
    fills = [{} for _ in rows]
    stops = ((233, 129, 129), (238, 238, 136), (128, 231, 127))
    for key in ('apo', 'qpay_conv'):
        values = [Decimal(str(row[key])) for row in rows if row.get(key) is not None]
        if not values:
            continue
        low, high = min(values), max(values)
        for row, fill in zip(rows, fills):
            if row.get(key) is None:
                continue
            position = (Decimal(str(row[key])) - low) / (high - low) if high > low else Decimal(0)
            segment = 0 if position <= Decimal('.5') else 1
            fraction = position * 2 - segment
            channels = [int((Decimal(a) + Decimal(b - a) * fraction).quantize(Decimal('1'), rounding=ROUND_HALF_UP))
                        for a, b in zip(stops[segment], stops[segment + 1])]
            fill[key] = '#' + ''.join(f'{channel:02x}' for channel in channels)
    return fills


def report_data(start, end):
    with connect() as conn, conn.transaction():
        conn.execute('SET TRANSACTION ISOLATION LEVEL REPEATABLE READ, READ ONLY')
        raw = conn.execute('''select p.*, (p.loaded_at<i.loaded_at or i.status='error') as stale
            from public.sales_performance p join public.sales_imports i using(source_id,report_date)
            where p.report_date between %s and %s order by p.report_date,p.source_id,p.store_id''', (start, end)).fetchall()
        imports = conn.execute('''select source_id, report_date,status,loaded_at,row_count,retained_count,error_code
            from public.sales_imports where report_date between %s and %s
            order by source_id,report_date''', (start, end)).fetchall()
        roster=active_roster(conn)
        catalog=source_catalog(conn)
    rows = restrict_to_roster(aggregate(raw),roster,catalog,imports,start,end)
    dealer = request.args.get('dealer', '')
    if dealer and dealer not in COLORS:
        raise PortalError('Unknown dealer.')
    relevant = [i for i in imports if not dealer or SOURCES[i['source_id']] == dealer]
    if dealer:
        rows = [r for r in rows if r['dealer'] == dealer]
    markets = sorted({r['market'] for r in rows})
    market, store_id = request.args.get('market', ''), request.args.get('store', '')
    if market:
        rows = [r for r in rows if r['market'] == market]
    stores = [{'id': r['dealer'] + ':' + r['store_id'], 'name': r['store']} for r in rows]
    if store_id:
        rows = [r for r in rows if r['dealer'] + ':' + r['store_id'] == store_id]
    total = ratios({k: sum((r[k] or 0) for r in rows) for k in METRICS})
    expected = sum(1 for s in SOURCES if not dealer or SOURCES[s] == dealer) * ((end - start).days + 1)
    complete = sum(i['status'] == 'complete' for i in relevant)
    times = [i['loaded_at'] for i in relevant if i['loaded_at']]
    return {'rows': rows, 'metric_fills': metric_fills(rows), 'totals': total, 'markets': markets, 'stores': stores,
            'dealers': [{'name': d, 'color': c} for d, c in COLORS.items()],
            'start': start.isoformat(), 'end': end.isoformat(), 'today': today().isoformat(),
            'calling_tree': {'active':bool(roster),'filename':roster[0]['filename'] if roster else None,
                             'stores':len(roster),'unmatched':sum(r.get('unmatched',False) for r in rows)},
            'coverage': {'complete': complete, 'expected': expected,
                         'retained': sum(i['retained_count'] for i in relevant),
                         'errors': sorted({SOURCES[i['source_id']] for i in relevant if i['status'] == 'error'})},
            'updated_at': max(times).isoformat() if times else None}


def workbook_bytes(data):
    fills = metric_fills(data['rows'])
    book = Workbook()
    sheet = book.active
    sheet.title = 'Sales Update'
    sheet.append([label for _, label in COLUMNS])
    for row in [*data['rows'], {'dealer': 'TOTAL', **data['totals']}]:
        values = [row.get(key, '') for key, _ in COLUMNS]
        sheet.append(values)
        # Remote store names are untrusted strings; never create spreadsheet formulas.
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = 's'
        for key, col in [('accessory', 9), ('apo', 10)]:
            sheet.cell(sheet.max_row, col).number_format = '$#,##0.00'
        sheet.cell(sheet.max_row, 13).number_format = '0.00"%"'
    dealers={r['dealer'] for r in data['rows']}
    accent=COLORS[next(iter(dealers))] if len(dealers)==1 else COLORS['Connect']
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor=accent.lstrip('#'))
    for idx in range(2,sheet.max_row):
        for cell in sheet[idx]:
            cell.font=Font(bold=True,color='083C51')
        for col in (9,10,11,13):
            value=sheet.cell(idx,col).value
            if value is None:
                continue
            color='B9E7F5' if col==11 else None
            if col==10:
                color=fills[idx-2].get('apo', '').lstrip('#')
            elif col==13:
                color=fills[idx-2].get('qpay_conv', '').lstrip('#')
            if color:
                sheet.cell(idx,col).fill=PatternFill('solid',fgColor=color)
    for cell in sheet[sheet.max_row]:
        cell.font=Font(bold=True,color='FFFFFF')
        cell.fill=PatternFill('solid',fgColor=accent.lstrip('#'))
    sheet.freeze_panes = 'D2'
    sheet.auto_filter.ref = sheet.dimensions
    for cells in sheet.columns:
        sheet.column_dimensions[cells[0].column_letter].width = min(45, max(14, max(len(str(c.value or '')) for c in cells) + 2))
    meta = book.create_sheet('Report details')
    meta.append(['Start', data['start']]); meta.append(['End', data['end']])
    meta.append(['Completed source-days', data['coverage']['complete']])
    meta.append(['Expected source-days', data['coverage']['expected']])
    meta.append(['Retained source rows', data['coverage']['retained']])
    refreshed = datetime.fromisoformat(data['updated_at']).astimezone(CENTRAL).strftime('%b %d, %Y %I:%M %p %Z') if data['updated_at'] else 'Not available'
    meta.append(['Latest refresh', refreshed])
    meta.append(['Calling Tree', data.get('calling_tree',{}).get('filename') or 'Not provided'])
    for row in meta:
        for cell in row:
            if isinstance(cell.value,str):
                cell.data_type='s'
    result = BytesIO(); book.save(result); result.seek(0)
    return result


def register_sales(app):
    @app.errorhandler(psycopg.Error)
    def sales_database_error(exc):
        app.logger.warning('Sales database unavailable (%s)', exc.sqlstate or 'connection_error')
        return jsonify(message='Sales data is temporarily unavailable. Please retry.'), 503

    @app.get('/api/internal/sales')
    @require_user()
    def sales_report():
        return jsonify(report_data(*date_range()))

    @app.get('/api/internal/sales/export')
    @require_user()
    def sales_export():
        start, end = date_range()
        return send_file(workbook_bytes(report_data(start, end)), as_attachment=True,
                         download_name=f'archet-sales-{start}-{end}.xlsx',
                         mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

    @app.post('/api/internal/sales/refresh')
    @require_user(admin=True)
    def sales_refresh():
        payload = request.get_json(silent=True) or {}
        if not isinstance(payload, dict) or type(payload.get('continue', False)) is not bool:
            raise PortalError('Invalid sync request.', 400)
        # Only the first batch resets freshness and error delays. Continuing must
        # not requeue successful downloads or repeatedly retry failing sources.
        return jsonify(run_batch(force_today=not payload.get('continue', False)))

    @app.post('/api/jobs/sales')
    def scheduled_sales():
        secret = os.getenv('CRON_SECRET', '')
        if not secret or not secrets.compare_digest(request.headers.get('Authorization', ''), 'Bearer ' + secret):
            return jsonify(message='Unauthorized'), 401
        if os.getenv('VERCEL_ENV') and os.getenv('VERCEL_ENV') != 'production':
            return jsonify(message='Production only'), 403
        result = run_batch()
        response = jsonify(result)
        response.headers['Cache-Control'] = 'no-store'
        return response
