"""Monthly quota reports, ported from spdi-ca and scoped to the active Calling Tree."""
import base64
import binascii
import calendar
from datetime import date
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from io import BytesIO
from pathlib import PureWindowsPath
import re
from zipfile import ZipFile, BadZipFile

from flask import g, jsonify, request, send_file
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException

try:
    from api import sales
    from api.calling_tree import active_roster, source_catalog
    from api.portal import PortalError, body, require_user
except ModuleNotFoundError:
    import sales
    from calling_tree import active_roster, source_catalog
    from portal import PortalError, body, require_user

GOALS = ('voice_goal', 'bts_goal', 'hsi_goal', 'accessory_goal', 'mim_goal')
MAX_BYTES = 2 * 1024 * 1024
ALIASES = {
    'market': ('market', 'market id'),
    'store': ('stores', 'store', 'store name', 'company'),
    'store_id': ('store id',),
    'voice_goal': ('voice', 'voice goal', 'voice goals'),
    'bts_goal': ('bts', 'bts goal', 'bts goals'),
    'hsi_goal': ('hsi', 'hint', 'hsi hint', 'hsi goal', 'hsi goals'),
    'accessory_goal': ('acc', 'acc goal', 'accessory', 'accessories', 'accessory goal'),
    'mim_goal': ('mim', 'mim goal', 'mim goals'),
}


def normalized(value):
    return ' '.join(str(value or '').strip().upper().split())


def month_value(value):
    if not isinstance(value, str) or not re.fullmatch(r'\d{4}-\d{2}', value):
        raise PortalError('Choose a valid goal month (YYYY-MM).')
    try:
        result = date.fromisoformat(value + '-01')
        if not 2000 <= result.year <= 2100:
            raise ValueError
        return result
    except ValueError:
        raise PortalError('Choose a goal month between 2000 and 2100.') from None


def goal_number(value, row):
    if value is None or value == '':
        return Decimal('0.00')
    try:
        text = str(value).strip().replace('$', '').replace(',', '')
        if text.startswith('(') and text.endswith(')'):
            text = '-' + text[1:-1]
        result = Decimal(text)
        if not result.is_finite() or not 0 <= result <= Decimal('999999999999.99'):
            raise ValueError
        return result.quantize(Decimal('.01'), rounding=ROUND_HALF_UP)
    except (InvalidOperation, ValueError):
        raise PortalError(f'Row {row}: goals must be nonnegative numbers.') from None


def parse_workbook(data, roster, dealer):
    """Accept legacy SPDI goals or the Store ID template; never execute workbook content."""
    if len(data) > MAX_BYTES:
        raise PortalError('Choose a workbook smaller than 2 MB.')
    book = cached = None
    try:
        with ZipFile(BytesIO(data)) as archive:
            if len(archive.infolist()) > 2000 or sum(i.file_size for i in archive.infolist()) > 25 * 1024 * 1024:
                raise PortalError('This workbook is too large when unpacked.')
        book = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
        cached = load_workbook(BytesIO(data), read_only=True, data_only=True, keep_links=False)
        candidates = []
        for sheet in book:
            if (sheet.max_column or 0) > 100 or (sheet.max_row or 0) > 5025:
                raise PortalError('Use at most 5,000 goal rows and 100 columns per worksheet.')
            for number, cells in enumerate(sheet.iter_rows(max_row=25), 1):
                headers = {re.sub(r'[^a-z0-9]+', ' ', str(c.value or '').lower()).strip(): i for i, c in enumerate(cells)}
                mapping = {key: next((headers[a] for a in names if a in headers), None) for key, names in ALIASES.items()}
                if all(mapping[k] is not None for k in ('market', 'store', *GOALS[:4])):
                    candidates.append((sheet, number, mapping))
                    break
        chosen = next((c for c in candidates if c[0].title.casefold() == 'goals'), candidates[0] if candidates else None)
        if not chosen:
            raise PortalError('Required columns: Market, Stores, Voice, BTS, HSI/HINT, Acc. MIM and Store ID are optional.')
        sheet, header, mapping = chosen
        stores = [s for s in roster if s['dealer'] == dealer]
        if not stores:
            raise PortalError('Upload a Calling Tree with stores for this dealer first.')
        by_id = {normalized(s['store_id']): s for s in stores}
        by_name = {}
        for store in stores:
            by_name.setdefault((normalized(store['market']), normalized(store['store'])), []).append(store)
        records = {}
        values = cached[sheet.title].iter_rows(min_row=header + 1)
        for number, (cells, cached_cells) in enumerate(zip(sheet.iter_rows(min_row=header + 1), values), header + 1):
            row = {}
            for key, col in mapping.items():
                cell = cells[col] if col is not None and col < len(cells) else None
                value = cell.value if cell else None
                if cell and cell.data_type == 'f':
                    value = cached_cells[col].value
                    if value is None:
                        raise PortalError(f'Row {number}: save calculated formulas in Excel or replace them with values.')
                if cell and cell.data_type == 'e':
                    raise PortalError(f'Row {number}: remove spreadsheet errors.')
                row[key] = value
            if all(v is None or v == '' for v in row.values()):
                continue
            if number - header > 5000:
                raise PortalError('Upload at most 5,000 goal rows.')
            sid = normalized(row['store_id'])
            matches = [by_id[sid]] if sid in by_id else []
            if not sid:
                matches = by_name.get((normalized(row['market']), normalized(row['store'])), [])
            if len(matches) != 1:
                raise PortalError(f'Row {number}: store does not uniquely match this dealer in the active Calling Tree. Use its Store ID from the template.')
            store = matches[0]
            key = store['store_id']
            record = records.setdefault(key, {k: store[k] for k in ('dealer', 'market', 'store', 'store_id')})
            for metric in GOALS:
                record[metric] = record.get(metric, Decimal(0)) + goal_number(row[metric], number)
                if record[metric] > Decimal('999999999999.99'):
                    raise PortalError(f'Row {number}: combined goal is too large.')
        if not records:
            raise PortalError('The worksheet contains no store goals.')
        return list(records.values())
    except (BadZipFile, InvalidFileException, OSError, ValueError, KeyError):
        raise PortalError('The file could not be read. Upload a valid .xlsx or .xlsm workbook.') from None
    finally:
        if book:
            book.close()
        if cached:
            cached.close()


def ratio(actual, goal):
    return actual / goal if goal else Decimal(0)


def summary_values(row, elapsed, days):
    actual, quota = Decimal(str(row['actual'])), Decimal(str(row['quota']))
    acc, goal = Decimal(str(row['acc_actual'])), Decimal(str(row['acc_goal']))
    trend = acc * days / elapsed if elapsed else Decimal(0)
    growth = ratio(actual, quota)
    achieved = ratio(trend, goal)
    return {**row, 'growth': growth, 'acc_remain': goal - acc,
            'per_day': (goal - acc) / (days - elapsed) if days > elapsed else Decimal(0),
            'trend': trend, 'achieved': achieved, 'overall': (growth + achieved) / 2}


def columns(*items):
    return [{'key': key, 'label': label, 'kind': kind} for key, label, kind in items]


LOCATION_COLUMNS = columns(('dealer', 'Dealer', 'text'), ('market', 'Market', 'text'), ('store', 'Store', 'text'))
SUMMARY_COLUMNS = columns(('rank', 'Rank', 'number')) + LOCATION_COLUMNS + columns(
    ('quota', 'Quota', 'number'), ('actual', 'Actual', 'number'), ('growth', 'Growth %', 'percent'),
    ('acc_goal', 'Acc Goal', 'money'), ('acc_actual', 'Acc Actual', 'money'), ('acc_remain', 'Acc Remain', 'money'),
    ('per_day', 'Per Day Goal', 'money'), ('trend', 'Acc Trend', 'money'),
    ('achieved', 'Achieved', 'percent'), ('overall', 'Overall Average', 'percent'))


def build_tables(goals, actuals, month, current):
    days = calendar.monthrange(month.year, month.month)[1]
    elapsed = max(0, min(days, (current - month).days + 1))
    by_store = {(r['dealer'], r['store_id']): r for r in actuals}
    tables = []
    summary = []
    metrics = {k: [] for k in ('voice', 'bts', 'hsi')}
    for goal in goals:
        source = by_store.get((goal['dealer'], goal['store_id']), {})
        location = {k: goal[k] for k in ('dealer', 'market', 'store', 'store_id')}
        missing = elapsed > 0 and source.get('new_activation') is None
        incomplete = elapsed > 0 and (missing or source.get('incomplete', False) or source.get('stale', False))
        values = {k: Decimal(str(source.get(k) or 0)) for k in ('new_activation', 'reactivation', 'bts', 'hsi', 'accessory')}
        values['voice'] = values['new_activation'] + values['reactivation'] - values['bts'] - values['hsi']
        for metric in metrics:
            target = Decimal(str(goal[metric + '_goal']))
            actual = values[metric]
            metrics[metric].append({**location, 'goal': target, 'actual': None if missing else actual,
                'remaining': None if missing else target - actual, 'growth': None if missing else ratio(actual, target),
                'incomplete': incomplete})
        row = summary_values({**location, 'quota': sum(Decimal(str(goal[k])) for k in GOALS if k != 'accessory_goal'),
            'actual': values['voice'] + values['bts'] + values['hsi'],
            'acc_goal': Decimal(str(goal['accessory_goal'])), 'acc_actual': values['accessory'],
            'incomplete': incomplete}, elapsed, days)
        if missing:
            for key in ('actual', 'growth', 'acc_actual', 'acc_remain', 'per_day', 'trend', 'achieved', 'overall'):
                row[key] = None
        summary.append(row)
    for metric, title in (('voice', 'VOICE GOALS'), ('bts', 'BTS GOALS'), ('hsi', 'HSI/HINT GOALS')):
        rows = sorted(metrics[metric], key=lambda r: (r['dealer'], r['market'], r['growth'] is None, -(r['growth'] or 0), r['store']))
        total = {'store': 'TOTAL', 'goal': sum(r['goal'] for r in rows),
                 'actual': sum(r['actual'] or 0 for r in rows), 'incomplete': any(r['incomplete'] for r in rows)}
        total.update(remaining=total['goal'] - total['actual'], growth=ratio(total['actual'], total['goal']))
        if rows and all(r['actual'] is None for r in rows):
            total.update(actual=None, remaining=None, growth=None)
        tables.append({'id': metric, 'title': title, 'columns': LOCATION_COLUMNS + columns(
            ('goal', metric.upper() + ' Goal', 'number'), ('actual', 'Actual', 'number'),
            ('remaining', 'Remaining', 'number'), ('growth', 'Growth %', 'percent')), 'rows': rows, 'total': total})
    summary.sort(key=lambda r: (r['dealer'], r['market'], r['overall'] is None, -(r['overall'] or 0), r['store']))
    group = None
    for row in summary:
        if group != (row['dealer'], row['market']):
            group = (row['dealer'], row['market']); position = rank = 0; previous = None
        position += 1
        if row['overall'] is not None:
            if previous != row['overall']:
                rank = position
            row['rank'] = rank
            previous = row['overall']
        else:
            row['rank'] = None
    total = summary_values({'store': 'TOTAL', **{k: sum(r[k] or 0 for r in summary)
        for k in ('quota', 'actual', 'acc_goal', 'acc_actual')},
        'incomplete': any(r['incomplete'] for r in summary)}, elapsed, days)
    if summary and all(r['actual'] is None for r in summary):
        for key in ('actual', 'growth', 'acc_actual', 'acc_remain', 'per_day', 'trend', 'achieved', 'overall'):
            total[key] = None
    tables.append({'id': 'summary', 'title': 'GOALS ACHIEVEMENT SUMMARY', 'columns': SUMMARY_COLUMNS,
                   'rows': summary, 'total': total})
    return tables, elapsed, days


def report_data():
    current = sales.today()
    dealer = request.args.get('dealer', '')
    if dealer and dealer not in sales.COLORS:
        raise PortalError('Unknown dealer.')
    with sales.connect() as conn, conn.transaction():
        conn.execute('SET TRANSACTION ISOLATION LEVEL REPEATABLE READ, READ ONLY')
        uploads = conn.execute('select * from public.quota_uploads order by goal_month desc,dealer').fetchall()
        month = month_value(request.args['month']) if request.args.get('month') else (uploads[0]['goal_month'] if uploads else current.replace(day=1))
        end = min(current, month.replace(day=calendar.monthrange(month.year, month.month)[1]))
        goals = conn.execute('''select q.*,s.market,s.store from public.quota_goals q
            join public.calling_tree_stores s on s.dealer=q.dealer and s.store_id=q.store_id
            join public.calling_tree_versions v on v.id=s.version_id and v.status='active'
            where q.goal_month=%s''', (month,)).fetchall()
        roster = active_roster(conn)
        raw = conn.execute('''select p.*, (p.loaded_at<i.loaded_at or i.status='error') as stale
            from public.sales_performance p join public.sales_imports i using(source_id,report_date)
            where p.report_date between %s and %s''', (month, end)).fetchall()
        imports = conn.execute('''select source_id,report_date,status,loaded_at from public.sales_imports
            where report_date between %s and %s''', (month, end)).fetchall()
        catalog = source_catalog(conn)
    actuals = sales.restrict_to_roster(sales.aggregate(raw), roster, catalog, imports, month, end)
    if dealer:
        goals = [g for g in goals if g['dealer'] == dealer]
    markets = sorted({g['market'] for g in goals})
    market = request.args.get('market', '')
    if market:
        goals = [g for g in goals if g['market'] == market]
    stores = [{'id': g['dealer'] + ':' + g['store_id'], 'name': g['store'], 'dealer': g['dealer']} for g in sorted(goals, key=lambda g: (g['store'],g['dealer']))]
    if request.args.get('store'):
        goals = [g for g in goals if g['dealer'] + ':' + g['store_id'] == request.args['store']]
    tables, elapsed, days = build_tables(goals, actuals, month, current)
    dealers = {g['dealer'] for g in goals}
    relevant = [i for i in imports if sales.SOURCES[i['source_id']] in dealers]
    times = [i['loaded_at'] for i in relevant if i['loaded_at']]
    return {'month': month.strftime('%Y-%m'), 'today': current.isoformat(), 'tables': tables,
            'months': sorted({u['goal_month'].strftime('%Y-%m') for u in uploads} | {month.strftime('%Y-%m')}, reverse=True),
            'uploads': [{'dealer': u['dealer'], 'filename': u['filename'], 'row_count': u['row_count'],
                         'uploaded_at': u['uploaded_at'].isoformat()} for u in uploads if u['goal_month'] == month and (not dealer or u['dealer'] == dealer)],
            'markets': markets, 'stores': stores, 'dealers': list(sales.COLORS), 'count': len(goals),
            'elapsed': elapsed, 'days': days, 'calling_tree_active': bool(roster),
            'incomplete': sum(r['incomplete'] for r in tables[0]['rows']),
            'updated_at': min(times).isoformat() if times else None}


def growth_color(value):
    if value is None:
        return None
    return next(color for threshold, color in ((1, '86EFAC'), (.75, 'BEF264'), (.5, 'FDE047'), (.25, 'FDBA74'), (-float('inf'), 'FCA5A5')) if value >= threshold)


def excel_file(book, filename):
    output = BytesIO()
    book.save(output); book.close(); output.seek(0)
    return send_file(output, as_attachment=True, download_name=filename,
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


def report_workbook(data):
    book = Workbook(); book.remove(book.active)
    dealers = {r['dealer'] for r in data['tables'][0]['rows']}
    accent = sales.COLORS[next(iter(dealers))] if len(dealers) == 1 else sales.COLORS['Connect']
    for table in data['tables']:
        sheet = book.create_sheet({'voice':'Voice Goals','bts':'BTS Goals','hsi':'HSI Goals','summary':'Achievement Summary'}[table['id']])
        cols = table['columns']
        sheet.append([table['title'] + ' | ' + data['month']])
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(cols))
        period = month_value(data['month'])
        last_day = period.replace(day=calendar.monthrange(period.year, period.month)[1]).isoformat()
        sheet.append(['Partial saved actuals; missing data is blank.' if data['incomplete'] else ('Month has not started.' if not data['elapsed'] else 'Saved actuals through ' + min(data['today'], last_day))])
        sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(cols))
        sheet.append([c['label'] for c in cols])
        for row in [*table['rows'], table['total']]:
            sheet.append([row.get(c['key'], '') for c in cols])
            for i, col in enumerate(cols, 1):
                cell = sheet.cell(sheet.max_row, i)
                if isinstance(cell.value, str):
                    cell.data_type = 's'
                cell.number_format = {'money':'$#,##0.00;($#,##0.00)', 'percent':'0.00%;(0.00%)', 'number':'#,##0.##;(#,##0.##)'}.get(col['kind'], 'General')
                fill = growth_color(row.get(col['key'])) if col['kind'] == 'percent' else None
                if fill:
                    cell.fill = PatternFill('solid', fgColor=fill)
        for n in (1, 3, sheet.max_row):
            for cell in sheet[n]:
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor=accent.lstrip('#'))
        for i, col in enumerate(cols, 1):
            sheet.column_dimensions[get_column_letter(i)].width = 36 if col['key'] == 'store' else 19
        sheet.freeze_panes = 'D4'
        sheet.auto_filter.ref = f'A3:{get_column_letter(len(cols))}{sheet.max_row-1}'
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = 'landscape'; sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
        sheet.page_setup.fitToWidth = 1; sheet.page_setup.fitToHeight = 0
    return book


def register_quota(app):
    @app.get('/api/internal/quota')
    @require_user()
    def quota_report():
        return jsonify(report_data())

    @app.get('/api/internal/quota/excel')
    @require_user()
    def quota_excel():
        data = report_data()
        return excel_file(report_workbook(data), 'Quota-Update-' + data['month'] + '.xlsx')

    @app.get('/api/internal/quota/template')
    @require_user(admin=True)
    def quota_template():
        dealer = request.args.get('dealer', '')
        if dealer not in sales.COLORS:
            raise PortalError('Choose a dealer for the template.')
        with sales.connect() as conn:
            roster = [s for s in active_roster(conn) if s['dealer'] == dealer]
        if not roster:
            raise PortalError('Upload a Calling Tree for this dealer first.')
        book = Workbook(); sheet = book.active; sheet.title = 'Goals'
        sheet.append(['Store ID', 'Market', 'Stores', 'Voice', 'BTS', 'HSI/HINT', 'Acc', 'MIM'])
        for row in roster:
            sheet.append([row['store_id'], row['market'], row['store'], 0, 0, 0, 0, 0])
            for cell in sheet[sheet.max_row][:3]:
                cell.data_type = 's'
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='FFFFFF'); cell.fill = PatternFill('solid', fgColor=sales.COLORS[dealer][1:])
        for col in 'ABCDEFGH':
            sheet.column_dimensions[col].width = 38 if col == 'C' else 18
        sheet.freeze_panes = 'D2'
        return excel_file(book, 'Quota-Goals-' + dealer + '.xlsx')

    @app.post('/api/internal/quota/upload')
    @require_user(admin=True)
    def quota_upload():
        payload = body()
        dealer = payload.get('dealer')
        if not isinstance(dealer, str) or dealer not in sales.COLORS:
            raise PortalError('Choose a dealer for these goals.')
        month = month_value(payload.get('month'))
        filename = payload.get('filename', '')
        if not isinstance(filename, str) or len(filename) > 240 or not filename.lower().endswith(('.xlsx', '.xlsm')):
            raise PortalError('Choose an .xlsx or .xlsm workbook.')
        filename = PureWindowsPath(filename).name
        content = payload.get('content', '')
        if not isinstance(content, str) or len(content) > (MAX_BYTES + 2) // 3 * 4:
            raise PortalError('Choose a workbook smaller than 2 MB.')
        try:
            data = base64.b64decode(content, validate=True)
        except (ValueError, binascii.Error):
            raise PortalError('The workbook upload is invalid.') from None
        apply = payload.get('apply', False)
        if not isinstance(apply, bool):
            raise PortalError('Invalid upload action.')
        with sales.connect() as conn, conn.transaction():
            # Shares activation's table lock so the roster cannot change during validation/save.
            conn.execute('LOCK TABLE public.calling_tree_versions IN SHARE MODE')
            roster = active_roster(conn)
            version = str(roster[0]['version_id']) if roster else ''
            if apply and payload.get('roster_version') != version:
                raise PortalError('The Calling Tree changed. Preview this workbook again.', 409)
            records = parse_workbook(data, roster, dealer)
            if apply:
                conn.execute('''insert into public.quota_uploads(goal_month,dealer,filename,row_count,uploaded_by)
                    values(%s,%s,%s,%s,%s) on conflict(goal_month,dealer) do update set
                    filename=excluded.filename,row_count=excluded.row_count,uploaded_by=excluded.uploaded_by,uploaded_at=now()''',
                    (month, dealer, filename, len(records), g.portal_user['id']))
                conn.execute('delete from public.quota_goals where goal_month=%s and dealer=%s', (month, dealer))
                with conn.cursor() as cursor:
                    cursor.executemany('''insert into public.quota_goals(goal_month,dealer,store_id,voice_goal,bts_goal,hsi_goal,accessory_goal,mim_goal)
                        values(%s,%s,%s,%s,%s,%s,%s,%s)''', [(month, dealer, r['store_id'], *(r[k] for k in GOALS)) for r in records])
        return jsonify(rows=records, count=len(records), roster_version=version,
                       message=f'{len(records)} store goals {"saved" if apply else "validated"} for {dealer}, {month:%B %Y}.')
