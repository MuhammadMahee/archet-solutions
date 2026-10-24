"""Store opening/closing reviews and inclusive RT-POS time-punch imports."""
from collections import Counter, defaultdict
from concurrent.futures import ThreadPoolExecutor
import csv
from decimal import Decimal, InvalidOperation
import base64
import binascii
from datetime import date, datetime, time as wall_time, timedelta
from hashlib import sha256
from io import BytesIO, StringIO
import json
import os
import re
import time
from urllib.parse import urlencode
from urllib.request import Request
from pathlib import PureWindowsPath
from uuid import UUID
from psycopg.types.json import Jsonb

from flask import g, jsonify, request, send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
import xlrd

from api.calling_tree import active_roster, clean
from api.portal import PortalError, body, require_user
from api.store_schedules import scheduled_hours
from api.store_hours_directory import (roster_with_schedules, hours_for_store, parse_schedule_workbook,
                                       map_schedule_rows, DEALER_LABELS, SOURCE_IDS)
from api.sales import (SOURCES, COLORS, MAX_REPORT_BYTES, ReportError, connect,
                       open_source, read_source, source_error, today)

REPORT_URL = 'https://myrtpos.com/newbdi/TimePunches.fwx'
LOCK_ID = 731043260927
COLUMNS = [('date', 'Date'), ('market', 'Market'), ('store', 'Store'),
           ('employee', 'Employee'), ('store_status', 'Opening Status'),
           ('opening_time', 'Opening Time'), ('camera_in', 'Camera In View'),
           ('rt_clockin', 'RT Clockin'), ('opening_status', 'Opening Status'),
           ('closing_time', 'Closing Time'), ('camera_out', 'Camera Out View'),
           ('rt_clockout', 'RT Clockout'), ('closing_status', 'Closing Status'),
           ('chargeback', 'Late Opening Chargeback')]
TIME_FIELDS = ('opening_time', 'closing_time', 'camera_in', 'camera_out')
OVERRIDE_FIELDS = {'date', 'market', 'store', 'employee', 'store_status', 'rt_clockin',
                   'opening_status', 'rt_clockout', 'closing_status', 'chargeback'}


def validate_cell_overrides(value):
    if not isinstance(value, dict) or set(value) - OVERRIDE_FIELDS:
        raise PortalError('Choose valid report columns to edit.')
    result = {}
    for field, raw in value.items():
        if field == 'date':
            try:
                day = date.fromisoformat(raw)
                if day > today():
                    raise ValueError
            except (TypeError, ValueError):
                raise PortalError('Enter a valid report date ending today or earlier.') from None
            result[field] = day.isoformat()
        elif field in ('rt_clockin', 'rt_clockout'):
            if raw is None or raw == '':
                result[field] = None
            else:
                try:
                    parsed = wall_time.fromisoformat(raw)
                    if parsed.tzinfo is not None or parsed.microsecond:
                        raise ValueError
                    result[field] = parsed.isoformat()
                except (TypeError, ValueError):
                    raise PortalError('Enter a valid RT clock time.') from None
        elif field == 'chargeback':
            if raw is None or raw == '':
                result[field] = None
                continue
            try:
                amount = Decimal(str(raw))
                if not amount.is_finite() or abs(amount) > 1000000 or amount != amount.quantize(Decimal('0.01')):
                    raise ValueError
            except (InvalidOperation, ValueError):
                raise PortalError('Enter a chargeback amount with up to two decimal places (maximum 1,000,000).') from None
            result[field] = float(amount)
        else:
            if not isinstance(raw, str) or len(raw) > 160 or any(ord(c) < 32 for c in raw):
                raise PortalError('Use plain text up to 160 characters in name, market, store, and status cells.')
            result[field] = raw.strip()
    return result


def parse_datetime(value):
    if isinstance(value, datetime):
        return value.replace(tzinfo=None)
    value = ' '.join(clean(value).split())
    if not value or value.replace('/', '').replace(':', '').strip() == '':
        return None
    for fmt in ('%m/%d/%Y %I:%M:%S %p', '%m/%d/%Y %I:%M %p',
                '%m/%d/%Y %H:%M:%S', '%m/%d/%Y %H:%M',
                '%Y-%m-%d %H:%M:%S', '%Y-%m-%dT%H:%M:%S'):
        try:
            return datetime.strptime(value, fmt)
        except ValueError:
            pass
    raise ReportError('invalid_punch_date')


def read_report(data):
    """RT-POS currently emits CSV despite btnExcel; also accept legacy XLS."""
    if data.startswith((b'\xd0\xcf\x11\xe0', b'\x09\x00', b'\x09\x02', b'\x09\x04', b'\x09\x08')):
        try:
            book = xlrd.open_workbook(file_contents=data)
            sheet = book.sheet_by_index(0)
            rows = []
            for n in range(sheet.nrows):
                rows.append([xlrd.xldate_as_datetime(c.value, book.datemode)
                             if c.ctype == xlrd.XL_CELL_DATE else c.value for c in sheet.row(n)])
        except (xlrd.XLRDError, IndexError, ValueError):
            raise ReportError('invalid_xls') from None
    else:
        text = data.decode('utf-8-sig', errors='replace')
        if '<html' in text[:2000].lower() or '<!doctype' in text[:2000].lower():
            raise ReportError('invalid_report_or_expired_cookies')
        delimiter = '\t' if '\t' in text.partition('\n')[0] else ','
        try:
            rows = list(csv.reader(StringIO(text), delimiter=delimiter, strict=True))
        except csv.Error:
            raise ReportError('invalid_csv') from None
    if not rows:
        raise ReportError('missing_headers')
    headers = {clean(v).lower(): i for i, v in enumerate(rows[0])}
    required = {'marketid', 'custno', 'company', 'name', 'lastname', 'clockintime', 'clockouttime'}
    if required - headers.keys():
        raise ReportError('report_columns_changed')
    result = {}
    for values in rows[1:]:
        if not any(clean(v) for v in values):
            continue
        if len(values) < max(headers[k] for k in required) + 1:
            raise ReportError('invalid_report_row')
        cell = lambda key: values[headers[key]] if key in headers else ''
        clock_in = parse_datetime(cell('clockintime'))
        if clock_in is None:
            # A missing date cannot advance the import watermark or be assigned to a day.
            raise ReportError('missing_clockin_date')
        clock_out = parse_datetime(cell('clockouttime'))
        if clock_out and clock_out < clock_in:
            raise ReportError('invalid_clockout_date')
        sid, store = clean(cell('custno')).upper(), clean(cell('company')).upper()
        if not sid and not store:
            raise ReportError('missing_store_identity')
        employee = ' '.join(filter(None, (clean(cell('name')), clean(cell('lastname'))))) or '-'
        key = sha256(json.dumps([sid, store, clean(cell('username')), employee, clock_in.isoformat()]).encode()).hexdigest()
        record = dict(punch_key=key, report_date=clock_in.date(), store_id=sid,
                      store=store, market=clean(cell('marketid')).upper(), employee=employee,
                      clock_in=clock_in, clock_out=clock_out)
        # Exact duplicate shifts collapse; a completed duplicate wins over an open one.
        if key not in result or (clock_out or datetime.min) > (result[key]['clock_out'] or datetime.min):
            result[key] = record
    return list(result.values())


def fetch_report(opener, start, end, deadline=None):
    payload = urlencode({'frmMarketID': '', 'frmStateID': '', 'frmRegionID': '',
                         'frmStoreType': '', 'frmStore': '', 'frmStart': start.strftime('%m/%d/%Y'),
                         'frmEnd': end.strftime('%m/%d/%Y'), 'btnExcel': 'click'}).encode()
    raw = read_source(opener, Request(REPORT_URL, data=payload, headers={
        'Content-Type': 'application/x-www-form-urlencoded', 'Referer': REPORT_URL}),
        MAX_REPORT_BYTES + 1, deadline=deadline)
    if len(raw) > MAX_REPORT_BYTES:
        raise ReportError('report_too_large')
    rows = read_report(raw)
    # RT-POS includes older, still-open shifts even when frmStart is later.
    # Assign days by clockintime and exclude carryovers from this replacement window.
    return [r for r in rows if start <= r['report_date'] <= end]


def sync_start(last_punch, last_complete, current):
    # Inclusive by design: refresh the last extracted clockintime date, including late clock-outs.
    return min(last_punch or last_complete or current.replace(day=1), current)


def seed_imports(conn, current, continuation=False):
    if continuation:
        return
    with conn.transaction():
        for source in SOURCES:
            punch = conn.execute('select max(report_date) as day from public.store_time_punches where source_id=%s', (source,)).fetchone()['day']
            complete = conn.execute("select max(report_date) as day from public.store_hours_imports where source_id=%s and status='complete'", (source,)).fetchone()['day']
            start = sync_start(punch, complete, current)
            conn.execute('''insert into public.store_hours_imports(source_id,report_date)
                select %s,d::date from generate_series(%s::date,%s::date,'1 day') d
                on conflict(source_id,report_date) do update set status='pending',error_code=null''', (source, start, current))
        conn.execute("update public.store_hours_imports set status='pending',error_code=null where status='error'")


def save_punches(conn, source, start, end, rows):
    with conn.transaction():
        conn.execute('delete from public.store_time_punches where source_id=%s and report_date between %s and %s', (source, start, end))
        if rows:
            fields = ('report_date', 'punch_key', 'store_id', 'market', 'store', 'employee', 'clock_in', 'clock_out')
            with conn.cursor() as cur:
                cur.executemany(f'''insert into public.store_time_punches(source_id,{','.join(fields)})
                    values({','.join(['%s'] * 9)})''', [(source, *(r[k] for k in fields)) for r in rows])
        counts = Counter(r['report_date'] for r in rows)
        with conn.cursor() as cur:
            cur.executemany('''update public.store_hours_imports set status='complete',loaded_at=now(),
                row_count=%s,error_code=null where source_id=%s and report_date=%s''',
                [(counts[start + timedelta(days=n)], source, start + timedelta(days=n)) for n in range((end-start).days+1)])


def source_worker(source, cookies):
    with connect() as conn:
        job = conn.execute("select min(report_date) as day from public.store_hours_imports where source_id=%s and status='pending'", (source,)).fetchone()['day']
        if not job:
            return {'source': source, 'completed': 0}
        end = min(job + timedelta(days=6), today())
        try:
            deadline = time.monotonic() + 200
            opener = open_source(cookies, deadline=deadline, report_url=REPORT_URL)
            rows = fetch_report(opener, job, end, deadline)
            save_punches(conn, source, job, end, rows)
            return {'source': source, 'completed': (end-job).days+1, 'rows': len(rows)}
        except (ReportError, OSError) as exc:
            code = source_error(exc)
            conn.execute("""update public.store_hours_imports set status='error',error_code=%s
                where source_id=%s and report_date between %s and %s""", (code, source, job, end))
            return {'source': source, 'completed': 0, 'error': code}


def run_sync(continuation=False):
    try:
        cookies = json.loads(os.getenv('RTPOS_COOKIES', '{}'))
        if not isinstance(cookies, dict):
            raise ValueError
    except ValueError:
        raise PortalError('RT-POS credentials need configuration.', 503) from None
    with connect() as conn:
        if not conn.execute('select pg_try_advisory_lock(%s) as locked', (LOCK_ID,)).fetchone()['locked']:
            return {'status': 'busy', 'remaining': 1, 'workers': []}
        try:
            seed_imports(conn, today(), continuation)
            with ThreadPoolExecutor(max_workers=6) as pool:
                workers = list(pool.map(lambda s: source_worker(s, cookies.get(s, {})), SOURCES))
            counts = conn.execute("""select count(*) filter(where status='pending') as pending,
                count(*) filter(where status='error') as failed from public.store_hours_imports""").fetchone()
            return dict(status='partial' if counts['pending'] or counts['failed'] else 'complete',
                        remaining=counts['pending'] + counts['failed'], workers=workers, **counts)
        finally:
            conn.execute('select pg_advisory_unlock(%s)', (LOCK_ID,))


def seconds(value):
    if not value:
        return None
    if isinstance(value, str):
        value = wall_time.fromisoformat(value)
    return value.hour * 3600 + value.minute * 60 + value.second


def evaluate(row):
    """Only camera evidence determines compliance; scheduled times are never guessed."""
    opening, closing, camera_in, camera_out = (seconds(row.get(k)) for k in TIME_FIELDS)
    result = {'store_status': 'Open' if camera_in is not None else 'Pending camera review',
              'opening_status': 'Pending camera review', 'closing_status': 'Pending camera review',
              'chargeback': None}
    if camera_in is not None:
        if opening is None:
            result['opening_status'] = 'Schedule required'
        else:
            delay = camera_in - opening
            result['opening_status'] = 'Late' if delay >= 60 else 'On Time'
            threshold = {'houston': 5, 'reno': 5, 'dallas': 10}.get(row['market'].strip().casefold())
            result['chargeback'] = -10 if row['dealer'] == 'Connect' and threshold and delay >= threshold * 60 else 0
    if camera_out is not None:
        if closing is None:
            result['closing_status'] = 'Schedule required'
        else:
            # Treat a close after midnight as the next day for an overnight schedule.
            if opening is not None and closing <= opening:
                closing += 86400
                if camera_out < opening:
                    camera_out += 86400
            result['closing_status'] = 'Early' if closing - camera_out >= 60 else 'On Time'
    return result


def iso_time(value):
    return value.isoformat() if value is not None else None


def employee_key(name):
    return 'person:' + ' '.join(name.split()).casefold() if name and name != '-' else ''


def build_rows(roster, punches, reviews, imports, start, end, manual_employees=()):
    by_id = defaultdict(list)
    by_name = defaultdict(list)
    name_counts = Counter((s['dealer'], s['store'].strip().upper()) for s in roster)
    for punch in punches:
        dealer = SOURCES[punch['source_id']]
        by_id[(dealer, punch['store_id'], punch['report_date'])].append(punch)
        by_name[(dealer, punch['store'].strip().upper(), punch['report_date'])].append(punch)
    saved = {(r['dealer'], r['store_id'], r['report_date'], r.get('employee_key', '')): r for r in reviews}
    reviewed_keys = defaultdict(set)
    for rd, rs, rday, key in saved:
        if key:
            reviewed_keys[rd, rs, rday].add(key)
    manual = defaultdict(dict)
    for entry in manual_employees:
        manual[entry['dealer'], entry['store_id'], entry['report_date']][entry['employee_key']] = entry['employee']
    loaded = {(r['source_id'], r['report_date']): r for r in imports}
    source_ids = {dealer: source for source, dealer in SOURCES.items()}
    rows = []
    for n in range((end-start).days+1):
        day = start + timedelta(days=n)
        for store in roster:
            dealer, sid = store['dealer'], store['store_id']
            source = store.get('source_id', source_ids.get(dealer))
            records = by_id.get((dealer, sid, day), []) if source else []
            name = store['store'].strip().upper()
            if source and not records and name_counts[dealer, name] == 1:
                records = by_name.get((dealer, name, day), [])
            grouped = defaultdict(list)
            for record in sorted(records, key=lambda r: (r['clock_in'], r['employee'])):
                grouped[employee_key(record['employee'])].append(record)
            manual_names = manual[dealer, sid, day]
            for key in sorted(manual_names):
                grouped.setdefault(key, [])
            # Preserve reviews whose imported employee temporarily disappears after a re-sync.
            for key in sorted(reviewed_keys[dealer, sid, day]):
                grouped.setdefault(key, [])
            if not grouped:
                grouped[''] = []
            schedule = hours_for_store(store, day)
            imported = loaded.get((source, day), {})
            for index, (key, shifts) in enumerate(grouped.items()):
                first = shifts[0] if shifts else None
                outs = [r['clock_out'] for r in shifts if r['clock_out']]
                review = saved.get((dealer, sid, day, key), {})
                # A legacy store-level review belongs only to the first employee, never all employees.
                if not review and index == 0:
                    review = saved.get((dealer, sid, day, ''), {})
                overrides = review.get('cell_overrides') or {}
                effective_day = date.fromisoformat(overrides.get('date', day.isoformat()))
                schedule = hours_for_store(store, effective_day)
                employee = first['employee'] if first else manual_names.get(key, key.removeprefix('person:') or '-')
                row = dict(date=day.isoformat(), dealer=dealer, market=store['market'], store=store['store'], store_id=sid,
                           employee=employee, employee_key=key, manual=key in manual_names,
                           rt_clockin=first['clock_in'].isoformat() if first else None,
                           rt_clockout=max(outs).isoformat() if outs else None,
                           revision=review.get('revision', 0), source_status=imported.get('status', 'not_synced') if source else 'manual_only',
                           schedule_period=schedule['period'],
                           scheduled_opening_time=iso_time(schedule.get('opening_time')),
                           scheduled_closing_time=iso_time(schedule.get('closing_time')),
                           **{k: iso_time(review.get(k) if review.get(k) is not None else schedule.get(k)) for k in TIME_FIELDS})
                row['source_date'] = day.isoformat()
                row['cell_overrides'] = overrides
                row['schedule_overrides'] = {f: review.get(f) is not None for f in ('opening_time', 'closing_time')}
                row['source_values'] = {k: row[k] for k in OVERRIDE_FIELDS if k in row}
                row.update({k: v for k, v in overrides.items() if k not in ('store_status','opening_status','closing_status','chargeback')})
                row.update(evaluate(row))
                row['source_values'].update({k: row[k] for k in ('store_status','opening_status','closing_status','chargeback')})
                row.update(overrides)
                rows.append(row)
                if len(rows) > 20000:
                    raise PortalError('Select a shorter date range or fewer stores (up to 20,000 employee rows).')
    return rows


def report_range():
    try:
        start = date.fromisoformat(request.args.get('start') or today().isoformat())
        end = date.fromisoformat(request.args.get('end') or start.isoformat())
    except ValueError:
        raise PortalError('Choose valid report dates.') from None
    if start > end or end > today() or (end-start).days > 366:
        raise PortalError('Choose up to 367 days ending today or earlier.')
    return start, end


def report_data(start, end):
    with connect() as conn:
        roster = roster_with_schedules(conn)
        reviews = conn.execute('select * from public.store_hours_reviews where report_date between %s and %s', (start, end)).fetchall()
        employee_reviews = conn.execute("""select * from public.store_hours_employee_reviews where
            report_date between %s and %s or cell_overrides->>'date' between %s and %s""",
            (start, end, start.isoformat(), end.isoformat())).fetchall()
        reviews += employee_reviews
        extra_days = sorted({r['report_date'] for r in employee_reviews if not start <= r['report_date'] <= end})
        args = (start, end, extra_days)
        punches = conn.execute('select * from public.store_time_punches where report_date between %s and %s or report_date=any(%s)', args).fetchall()
        manual = conn.execute('select * from public.store_hours_manual_employees where report_date between %s and %s or report_date=any(%s)', args).fetchall()
        imports = conn.execute('select * from public.store_hours_imports where report_date between %s and %s or report_date=any(%s)', args).fetchall()
        earliest = conn.execute('select min(report_date) as day from public.store_hours_imports').fetchone()['day']
    dealer, market, sid = (request.args.get(k, '') for k in ('dealer', 'market', 'store'))
    candidates = [s for s in roster if not dealer or s['dealer'] == dealer]
    identities = {(s['dealer'],s['store_id']) for s in candidates}
    corrections = [r for r in employee_reviews if (r['dealer'],r['store_id']) in identities and
                   start.isoformat() <= (r.get('cell_overrides') or {}).get('date',r['report_date'].isoformat()) <= end.isoformat()]
    markets = sorted({s['market'] for s in candidates} |
                     {r['cell_overrides']['market'] for r in corrections if 'market' in (r.get('cell_overrides') or {})})
    corrected_markets = {(r['dealer'],r['store_id']) for r in corrections if (r.get('cell_overrides') or {}).get('market') == market}
    candidates = [s for s in candidates if not market or s['market']==market or (s['dealer'],s['store_id']) in corrected_markets]
    store_names = {s['dealer']+':'+s['store_id']:s['store'] for s in candidates}
    for review in corrections:
        identity = review['dealer']+':'+review['store_id']
        if identity in store_names and 'store' in (review.get('cell_overrides') or {}):
            store_names[identity] = review['cell_overrides']['store']
    stores = [{'id':key,'name':name} for key,name in sorted(store_names.items(),key=lambda pair:pair[1])]
    selected = [s for s in candidates if not sid or s['dealer']+':'+s['store_id']==sid]
    if len(selected) * ((end-start).days+1) > 20000:
        raise PortalError('Select a dealer, market, or shorter date range (up to 20,000 rows).')
    rows = build_rows(selected, punches, reviews, imports, start, end, manual)
    for day in extra_days:
        identities = {(r['dealer'],r['store_id']) for r in employee_reviews if r['report_date']==day}
        moved_stores = [s for s in selected if (s['dealer'],s['store_id']) in identities]
        rows += build_rows(moved_stores, punches, reviews, imports, day, day, manual)
    rows = [r for r in rows if start.isoformat() <= r['date'] <= end.isoformat()]
    occupied = {(r['dealer'],r['store_id'],r['date']) for r in rows if r['employee_key']}
    rows = [r for r in rows if r['employee_key'] or r['revision'] or r['rt_clockin'] or r['rt_clockout'] or
            (r['dealer'],r['store_id'],r['date']) not in occupied]
    rows.sort(key=lambda r:(r['date'],r['dealer'],r['market'],r['store'],r['employee'],r['source_date']))
    if len(rows)>20000:
        raise PortalError('Select a shorter date range (up to 20,000 rows).')
    rows = [r for r in rows if not market or r['market']==market]
    included = {(r['dealer'],r['store_id']) for r in rows}
    selected = [s for s in selected if (s['dealer'],s['store_id']) in included]
    query = request.args.get('search', '').strip().casefold()
    if query:
        rows = [r for r in rows if query in ' '.join(r[k] for k in ('employee','store','dealer','market')).casefold()]
    relevant_sources = {r.get('source_id') for r in selected if r.get('source_id')}
    relevant = [i for i in imports if i['source_id'] in relevant_sources and start <= i['report_date'] <= end]
    updated = max((i['loaded_at'] for i in relevant if i['loaded_at']), default=None)
    return dict(rows=rows, start=start.isoformat(), end=end.isoformat(), today=today().isoformat(),
                earliest=(earliest or today().replace(day=1)).isoformat(),
                dealers=[{'id': d, 'name': DEALER_LABELS.get(d, d)} for d in sorted({s['dealer'] for s in roster})],
                markets=markets, stores=stores,
                directory=[{k: s[k] for k in ('dealer','store_id','market','store')} for s in roster],
                manual_store_count=sum(not s.get('source_id') for s in selected),
                updated_at=updated.isoformat() if updated else None,
                coverage={'complete': sum(i['status'] == 'complete' for i in relevant),
                          'expected': len(relevant_sources) * ((end-start).days+1),
                          'errors': sorted({SOURCES[i['source_id']] for i in relevant if i['status'] == 'error'})})


def display_value(key, value):
    if key == 'chargeback':
        return 'Pending' if value is None else ('-' if value < 0 else '') + '$' + format(abs(value), '.2f').removesuffix('.00')
    if key in TIME_FIELDS or key in ('rt_clockin', 'rt_clockout'):
        if not value:
            return {'rt_clockin': 'No Clock In', 'rt_clockout': 'No Clock Out'}.get(key, '-')
        value = datetime.fromisoformat(value).time() if 'T' in value else wall_time.fromisoformat(value)
        return value.strftime('%I:%M:%S %p').lstrip('0')
    return value


def excel_report(data):
    book = Workbook()
    sheet = book.active
    sheet.title = 'Store Opening Closing'
    sheet.append([f"STORE OPENING / CLOSING | {data['start']} to {data['end']}"])
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(COLUMNS))
    sheet.append([heading for _, heading in COLUMNS])
    for row in data['rows']:
        sheet.append([row[key] if key == 'chargeback' and row[key] is not None else display_value(key, row[key]) for key, _ in COLUMNS])
        for col, (key, _) in enumerate(COLUMNS, 1):
            cell = sheet.cell(sheet.max_row, col)
            if isinstance(cell.value, str):
                cell.data_type = 's'  # Imported names are never executable Excel formulas.
            if key == 'chargeback' and isinstance(cell.value, (int, float)):
                cell.number_format = '$0.00;[Red]-$0.00'
            if row[key] in ('Late', 'Early') or (key == 'chargeback' and row[key] is not None and row[key] < 0):
                cell.fill = PatternFill('solid', fgColor='FFC7CE')
            elif row[key] in ('Open', 'On Time'):
                cell.fill = PatternFill('solid', fgColor='C6EFCE')
            if key in ('rt_clockin', 'rt_clockout') and not row[key]:
                cell.font = Font(color='C00000')
    for row in sheet.iter_rows(min_row=1, max_row=2):
        for cell in row:
            cell.fill = PatternFill('solid', fgColor='203C66')
            cell.font = Font(color='FFFFFF', bold=True)
            cell.alignment = Alignment(horizontal='center')
    widths = [14, 20, 34, 30, 27, 17, 18, 18, 27, 17, 18, 18, 27, 25]
    for n, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(n)].width = width
    sheet.freeze_panes = 'E3'
    sheet.auto_filter.ref = f'A2:N{sheet.max_row}'
    meta = book.create_sheet('Report details')
    for values in [('Source', REPORT_URL), ('From', data['start']), ('To', data['end']),
                   ('Latest sync', data['updated_at'] or 'Not synced'),
                   ('Source coverage', f"{data['coverage']['complete']}/{data['coverage']['expected']} account-days"),
                   ('Rules', 'Late: 1 minute after opening. Early: 1 minute before closing.'),
                   ('Connect chargebacks', 'Houston/Reno: 5 minutes; Dallas: 10 minutes. Charge: -$10.'),
                   ('Manual corrections', 'Every form column supports saved corrections. Manually edited statuses and chargebacks override automatic calculations until reset to source/default.'),
                   ('Pending', 'Camera review and scheduled hours drive automatic compliance. RT punches do not determine compliance.'),
                   ('RT aggregation', 'One row per unique employee/store/date. Earliest clock-in and latest clock-out for that employee.'),
                   ('Chargeback scope', 'Assessed separately for each late employee, using their camera-in time.'),
                   ('Schedules', 'OPENING TIMING FOR ALL DEALERS.xlsx: Monday-Friday, Saturday, Sunday; date-specific overrides take priority.'),
                   ('Schedule correction', 'Connect Monday-Saturday opening is 10:00 AM, confirmed by the user.'),
                   ('Times', 'Source and camera times are store-local wall times.')]:
        meta.append(values)
    meta.column_dimensions['A'].width = 26
    meta.column_dimensions['B'].width = 110
    output = BytesIO()
    book.save(output)
    output.seek(0)
    return output


def store_day_rows(conn, store, day):
    punches = conn.execute('select * from public.store_time_punches where report_date=%s', (day,)).fetchall()
    args = (store['dealer'], store['store_id'], day)
    reviews = conn.execute('select * from public.store_hours_reviews where dealer=%s and store_id=%s and report_date=%s', args).fetchall()
    reviews += conn.execute('select * from public.store_hours_employee_reviews where dealer=%s and store_id=%s and report_date=%s', args).fetchall()
    manual = conn.execute('select * from public.store_hours_manual_employees where dealer=%s and store_id=%s and report_date=%s', args).fetchall()
    return build_rows([store], punches, reviews, [], day, day, manual)


def review_identity(payload):
    dealer, sid = payload.get('dealer'), payload.get('store_id')
    try:
        day = date.fromisoformat(payload.get('date', ''))
    except (ValueError, TypeError):
        raise PortalError('Choose a valid report date.') from None
    if day > today() or not all(isinstance(v, str) and 0 < len(v) <= 160 for v in (dealer, sid)):
        raise PortalError('Choose a valid store and date.')
    return dealer, sid, day


def lock_store_day(conn, dealer, sid, day):
    conn.execute('select pg_advisory_xact_lock(hashtextextended(%s,0))',
                 (json.dumps([dealer, sid, day.isoformat()]),))


def save_review(payload, user_id):
    dealer, sid, day = review_identity({**payload, 'date': payload.get('source_date', payload.get('date'))})
    overrides = validate_cell_overrides(payload['cell_overrides']) if 'cell_overrides' in payload else None
    key, revision = payload.get('employee_key', ''), payload.get('revision', 0)
    if not isinstance(key, str) or len(key) > 500 or type(revision) is not int or revision < 0:
        raise PortalError('Reload this review before editing.')
    times = {}
    for field in TIME_FIELDS:
        value = payload.get(field)
        if value is None or value == '':
            times[field] = None
        elif isinstance(value, str) and re.fullmatch(r'\d{2}:\d{2}(:\d{2})?', value):
            try:
                times[field] = wall_time.fromisoformat(value)
            except ValueError:
                raise PortalError('Enter valid opening, closing, and camera times.') from None
        else:
            raise PortalError('Enter valid opening, closing, and camera times.')
    with connect() as conn:
        with conn.transaction():
            lock_store_day(conn, dealer, sid, day)
            store = next((s for s in roster_with_schedules(conn) if s['dealer'] == dealer and s['store_id'] == sid), None)
            if store is None:
                raise PortalError('This store is no longer in the directory.', 409)
            row = next((r for r in store_day_rows(conn, store, day) if r['employee_key'] == key), None)
            if row is None:
                raise PortalError('This employee row has changed. Reload before saving.', 409)
            if row['revision'] != revision:
                raise PortalError('Someone updated this review. Discard changes and reload before saving again.', 409)
            if overrides is None:
                overrides = row.get('cell_overrides') or {}
            effective_day = date.fromisoformat(overrides.get('date', day.isoformat()))
            schedule = hours_for_store(store, effective_day)
            for field in ('opening_time', 'closing_time'):
                unchanged_default = not row.get('schedule_overrides', {}).get(field) and iso_time(times[field]) == row.get(field)
                if times[field] == schedule.get(field) or unchanged_default:
                    times[field] = None
            conn.execute("select set_config('archet.store_hours_employee',%s,true)", (row.get('employee', ''),))
            result = conn.execute('''insert into public.store_hours_employee_reviews
                (dealer,store_id,report_date,employee_key,opening_time,closing_time,camera_in,camera_out,updated_by,revision,cell_overrides)
                values(%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)
                on conflict(dealer,store_id,report_date,employee_key) do update set
                opening_time=excluded.opening_time,closing_time=excluded.closing_time,
                camera_in=excluded.camera_in,camera_out=excluded.camera_out,
                updated_by=excluded.updated_by,updated_at=now(),revision=excluded.revision,
                cell_overrides=excluded.cell_overrides returning revision''',
                (dealer, sid, day, key, *(times[k] for k in TIME_FIELDS), user_id, revision + 1, Jsonb(overrides))).fetchone()
    return result['revision']


def add_manual_employee(payload, user_id):
    dealer, sid, day = review_identity(payload)
    name = payload.get('employee')
    if not isinstance(name, str) or not name.strip() or len(name) > 160 or name.strip() == '-':
        raise PortalError('Enter an employee name (up to 160 characters).')
    name = ' '.join(name.split())
    key = employee_key(name)
    with connect() as conn:
        with conn.transaction():
            lock_store_day(conn, dealer, sid, day)
            store = next((s for s in roster_with_schedules(conn) if s['dealer'] == dealer and s['store_id'] == sid), None)
            if not store:
                raise PortalError('Choose a store from the directory.')
            rows = store_day_rows(conn, store, day)
            if any(r['employee_key'] == key for r in rows):
                raise PortalError('This employee already has a row for this store and date.', 409)
            if len(rows) >= 100:
                raise PortalError('A store can have up to 100 employee rows per date.')
            conn.execute('''insert into public.store_hours_manual_employees
                (dealer,store_id,report_date,employee_key,employee,created_by) values(%s,%s,%s,%s,%s,%s)''',
                (dealer, sid, day, key, name, user_id))
    return key


def preview_schedule_upload(payload, user_id):
    filename, content = payload.get('filename'), payload.get('content')
    if not isinstance(filename, str) or len(filename) > 250 or not filename.lower().endswith('.xlsx'):
        raise PortalError('Choose an .xlsx schedule workbook.')
    if not isinstance(content, str) or len(content) > 2800000:
        raise PortalError('Choose a workbook smaller than 2 MB.')
    try:
        raw = base64.b64decode(content, validate=True)
    except (ValueError, binascii.Error):
        raise PortalError('The upload was incomplete. Choose the workbook again.') from None
    parsed = parse_schedule_workbook(raw)
    with connect() as conn:
        rows = map_schedule_rows(parsed, roster_with_schedules(conn))
        conn.execute("delete from public.store_hours_uploads where applied_at is null and uploaded_at<now()-interval '1 day'")
        preview = conn.execute('''insert into public.store_hours_uploads(filename,rows,uploaded_by)
            values(%s,%s,%s) returning id''', (PureWindowsPath(filename).name, Jsonb(rows), user_id)).fetchone()['id']
    return dict(preview_id=str(preview), rows=rows, added=sum(r['action'] == 'Add' for r in rows),
                updated=sum(r['action'] == 'Update' for r in rows), corrected=sum(bool(r.get('corrected_connect_opening')) for r in rows))


def apply_schedule_upload(payload, user_id):
    try:
        preview = str(UUID(str(payload.get('preview_id', ''))))
    except ValueError:
        raise PortalError('Preview the workbook before applying it.') from None
    with connect() as conn:
        with conn.transaction():
            conn.execute('select pg_advisory_xact_lock(731043260928)')
            upload = conn.execute('''select * from public.store_hours_uploads where id=%s and uploaded_by=%s
                and (uploaded_at>now()-interval '1 day' or applied_at is not null) for update''', (preview, user_id)).fetchone()
            if not upload:
                raise PortalError('This preview expired. Preview the workbook again.', 409)
            if upload['applied_at']:
                return len(upload['rows'])
            # Revalidate identity before writing in case another upload added stores since preview.
            remapped = map_schedule_rows(upload['rows'], roster_with_schedules(conn))
            if [(r['dealer'],r['store_id']) for r in remapped] != [(r['dealer'],r['store_id']) for r in upload['rows']]:
                raise PortalError('The directory changed. Preview this workbook again.', 409)
            conn.execute("select set_config('archet.store_hours_actor',%s,true)", (str(user_id),))
            with conn.cursor() as cur:
                cur.executemany('''insert into public.store_hours_directory
                    (dealer,store_id,market,store,company,carrier,source_id,schedules) values(%s,%s,%s,%s,%s,%s,%s,%s)
                    on conflict(dealer,store_id) do update set market=excluded.market,store=excluded.store,
                    company=excluded.company,carrier=excluded.carrier,schedules=excluded.schedules,updated_at=now()''',
                    [(r['dealer'],r['store_id'],r['market'],r['store'],r['company'],r['carrier'],r['source_id'],Jsonb(r['schedules'])) for r in remapped])
            conn.execute('update public.store_hours_uploads set applied_at=now() where id=%s', (preview,))
    return len(remapped)


def version_history(start, end):
    """Read-only, keyset-paginated audit history; never expose a mutation endpoint."""
    before = request.args.get('before', '')
    if before and (not before.isascii() or not before.isdigit() or len(before) > 19 or not 0 < int(before) < 2**63):
        raise PortalError('Choose a valid history page.')
    clauses = ["""(report_date between %s and %s or report_date is null or
        before_data->'cell_overrides'->>'date' between %s and %s or
        after_data->'cell_overrides'->>'date' between %s and %s)"""]
    args = [start, end, start.isoformat(), end.isoformat(), start.isoformat(), end.isoformat()]
    for parameter, column in [('dealer', 'dealer'), ('market', 'market')]:
        value = request.args.get(parameter, '')
        if value:
            if column == 'market':
                clauses.append("(market=%s or before_data->'cell_overrides'->>'market'=%s or after_data->'cell_overrides'->>'market'=%s)")
                args.extend([value] * 3)
            else:
                clauses.append(column + '=%s')
                args.append(value)
    store = request.args.get('store', '')
    if store:
        parts = store.split(':', 1)
        if len(parts) != 2 or not all(parts):
            raise PortalError('Choose a valid store.')
        clauses.append('dealer=%s and store_id=%s')
        args.extend(parts)
    search = request.args.get('search', '').strip()
    if search:
        clauses.append("strpos(lower(concat_ws(' ',employee,store,dealer,market,actor_name,"
                       "before_data->'punches',after_data->'punches',"
                       "before_data->'cell_overrides',after_data->'cell_overrides')),lower(%s))>0")
        args.append(search)
    if before:
        clauses.append('id<%s')
        args.append(int(before))
    with connect() as conn:
        rows = conn.execute('select * from public.store_hours_versions where ' + ' and '.join(clauses) +
                            ' order by id desc limit 51', args).fetchall()
        tracking = conn.execute('select min(recorded_at) as started from public.store_hours_versions').fetchone()['started']
    result = []
    for row in rows[:50]:
        result.append({**row, 'id': str(row['id']), 'transaction_id': str(row['transaction_id']),
                       'recorded_at': row['recorded_at'].isoformat(),
                       'report_date': row['report_date'].isoformat() if row['report_date'] else None})
    return dict(versions=result, next_cursor=str(rows[49]['id']) if len(rows) > 50 else None,
                tracking_started=tracking.isoformat() if tracking else None)


def register_store_hours(app):
    @app.get('/api/internal/store-hours')
    @require_user()
    def get_report():
        return jsonify(report_data(*report_range()))

    @app.get('/api/internal/store-hours/history')
    @require_user()
    def get_version_history():
        return jsonify(version_history(*report_range()))

    @app.post('/api/internal/store-hours/refresh')
    @require_user(admin=True)
    def sync_report():
        return jsonify(run_sync(body().get('continue') is True))

    @app.post('/api/internal/store-hours/review')
    @require_user()
    def review_report():
        return jsonify(status='saved', revision=save_review(body(), g.portal_user['id']))

    @app.post('/api/internal/store-hours/employees')
    @require_user()
    def create_employee():
        return jsonify(status='saved', employee_key=add_manual_employee(body(), g.portal_user['id']))

    @app.post('/api/internal/store-hours/schedules/preview')
    @require_user(admin=True)
    def preview_schedules():
        return jsonify(preview_schedule_upload(body(), g.portal_user['id']))

    @app.post('/api/internal/store-hours/schedules/apply')
    @require_user(admin=True)
    def apply_schedules():
        return jsonify(status='saved', count=apply_schedule_upload(body(), g.portal_user['id']))

    @app.get('/api/internal/store-hours/export')
    @require_user()
    def export_report():
        data = report_data(*report_range())
        return send_file(excel_report(data), mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                         as_attachment=True, download_name=f"store-opening-closing-{data['start']}-{data['end']}.xlsx")
