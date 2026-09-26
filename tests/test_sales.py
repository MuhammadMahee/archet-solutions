from datetime import date
from decimal import Decimal
from io import BytesIO
from threading import Barrier
from contextlib import nullcontext
from http.client import IncompleteRead
from urllib.error import HTTPError, URLError
import ssl

import pytest
from openpyxl import load_workbook

from api import index, sales
from test_portal import store, client, login, write


def row(source='arm1', day=date(2026, 9, 1), sid='S1', **changes):
    return {'source_id': source, 'report_date': day, 'store_id': sid,
            'market': 'MARKET', 'store': 'STORE', 'stale': False,
            **{k: 0 for k in sales.METRICS}, **changes}


def test_arm_accounts_with_same_store_id_remain_separate():
    rows = [row(new_activation=2, total_boxes=2, accessory=Decimal('50'), qpay=4),
            row('arm2', new_activation=99, total_boxes=99),
            row('connect', new_activation=3, total_boxes=3),
            row(day=date(2026, 9, 2), total_boxes=1, accessory=Decimal('10'), qpay=2)]
    result = sales.aggregate(rows)
    arm = next(r for r in result if r['dealer'] == 'AMQ')
    assert arm['total_boxes'] == 3
    assert arm['apo'] == Decimal('20.00')
    assert arm['qpay_conv'] == Decimal('50.00')
    assert next(r for r in result if r['dealer'] == 'ARM')['total_boxes'] == 99
    assert len(result) == 3


def test_arm_freshness_is_independent_per_account():
    result = sales.aggregate([row(stale=True, total_boxes=5), row('arm2', total_boxes=3)])
    by_dealer = {r['dealer']: r for r in result}
    assert by_dealer['AMQ']['total_boxes'] == 5
    assert by_dealer['AMQ']['stale'] is True
    assert by_dealer['ARM']['total_boxes'] == 3
    assert by_dealer['ARM']['stale'] is False


def test_fresh_copy_wins_within_one_dealer():
    result = sales.aggregate([row(stale=True, total_boxes=5), row(total_boxes=3)])
    assert result[0]['total_boxes'] == 3
    assert result[0]['stale'] is False


def test_zero_denominators_and_half_up_rounding():
    assert sales.ratios({'accessory': Decimal('1.005'), 'total_boxes': 1, 'qpay': 0})['apo'] == Decimal('1.01')
    assert sales.ratios({'accessory': 0, 'total_boxes': 0, 'qpay': 0})['qpay_conv'] == 0


def test_metric_gradients_scale_columns_independently_and_continuously():
    rows = [{'apo': n * 5, 'qpay_conv': 200 - n * 50} for n in range(5)]
    fills = sales.metric_fills(rows)
    assert fills[0] == {'apo': '#e98181', 'qpay_conv': '#80e77f'}
    assert fills[2] == {'apo': '#eeee88', 'qpay_conv': '#eeee88'}
    assert fills[4] == {'apo': '#80e77f', 'qpay_conv': '#e98181'}
    assert len({fill['apo'] for fill in fills}) == 5
    assert sales.metric_fills(rows[1:4])[0]['apo'] == '#e98181'


def test_metric_gradients_handle_missing_and_equal_values():
    assert sales.metric_fills([]) == []
    assert sales.metric_fills([{'apo': None, 'qpay_conv': None}]) == [{}]
    assert sales.metric_fills([{'apo': 0, 'qpay_conv': 0}] * 2) == [
        {'apo': '#e98181', 'qpay_conv': '#e98181'}] * 2
    assert sales.metric_fills([{'apo': 10, 'qpay_conv': None}]) == [{'apo': '#e98181'}]


def test_excel_uses_the_same_metric_gradients_as_the_report():
    rows = sales.aggregate([row(sid=str(n), total_boxes=n, accessory=n*n*5, qpay=4) for n in range(5)])
    rows.append({**rows[-1], 'apo': None, 'qpay_conv': None})
    data = {'rows': rows, 'totals': rows[0], 'start': '2026-09-01', 'end': '2026-09-01',
            'coverage': {'complete': 6, 'expected': 6, 'retained': 0}, 'updated_at': None}
    sheet = load_workbook(sales.workbook_bytes(data)).active
    for index, fill in enumerate(sales.metric_fills(rows), 2):
        for key, column in [('apo', 10), ('qpay_conv', 13)]:
            cell = sheet.cell(index, column)
            if key in fill:
                assert cell.fill.fgColor.rgb[-6:].lower() == fill[key][1:]
            else:
                assert cell.fill.patternType is None


@pytest.mark.parametrize('value', ['NaN', 'Infinity', 'wrong'])
def test_bad_numeric_values_reject_whole_report(value):
    with pytest.raises(sales.ReportError):
        sales.number(value)


def test_new_account_sessions_discard_shared_fw_cookie(monkeypatch):
    captured = []
    class Response:
        def __enter__(self): return self
        def __exit__(self, *args): pass
        def read(self, *args): return b'<input name="frmStart"><input name="frmEnd">'
    class Opener:
        def open(self, *args, **kwargs): return Response()
    def opener(*handlers):
        captured.extend(list(handlers[1].cookiejar))
        return Opener()
    monkeypatch.setattr(sales, 'build_opener', opener)
    sales.open_source({'FW_SessionID':'SHARED', 'sec85952EAF_id':'user', 'sec85952EAF_pd':'secret'})
    assert {c.name for c in captured} == {'sec85952EAF_id', 'sec85952EAF_pd'}
    assert all(c.secure and c.domain == 'myrtpos.com' for c in captured)


def test_redirect_cannot_send_cookie_to_unrelated_host():
    with pytest.raises(sales.ReportError):
        sales.SameHostRedirect().redirect_request(None, None, 302, '', {}, 'https://example.com/')


@pytest.mark.parametrize('error', [TimeoutError(), URLError(TimeoutError()),
    ConnectionResetError(), IncompleteRead(b'partial'),
    HTTPError(sales.REPORT_URL, 503, 'Unavailable', {}, None)])
def test_source_transport_retries_and_recovers(error, monkeypatch):
    attempts, delays = [], []
    class Response:
        def __enter__(self): return self
        def __exit__(self, *args): pass
        def read(self, limit): return b'complete'
    class Opener:
        def open(self, req, timeout):
            attempts.append(timeout)
            if len(attempts) < 3: raise error
            return Response()
    monkeypatch.setattr(sales.time, 'sleep', delays.append)
    assert sales.read_source(Opener(), sales.REPORT_URL, 100) == b'complete'
    assert len(attempts) == 3 and delays == [1, 3]


@pytest.mark.parametrize('error,code', [
    (HTTPError(sales.REPORT_URL, 403, 'Forbidden', {}, None), 'source_http_403'),
    (URLError(ssl.SSLCertVerificationError()), 'source_tls_error')])
def test_source_does_not_retry_denied_access_or_bad_certificates(error, code, monkeypatch):
    class Opener:
        def open(self, *args, **kwargs): raise error
    monkeypatch.setattr(sales.time, 'sleep', lambda _: pytest.fail('must not retry'))
    with pytest.raises(type(error)):
        sales.read_source(Opener(), sales.REPORT_URL, 100)
    assert sales.source_error(error) == code


def test_source_retries_stop_at_attempt_limit(monkeypatch):
    calls = []
    class Opener:
        def open(self, *args, **kwargs):
            calls.append(1)
            raise URLError(TimeoutError())
    monkeypatch.setattr(sales.time, 'sleep', lambda _: None)
    with pytest.raises(URLError) as error:
        sales.read_source(Opener(), sales.REPORT_URL, 100)
    assert len(calls) == 3
    assert sales.source_error(error.value) == 'source_timeout'


def test_source_deadline_caps_timeout_and_stops_retries(monkeypatch):
    monkeypatch.setattr(sales.time, 'monotonic', lambda: 100)
    monkeypatch.setattr(sales.time, 'sleep', lambda _: pytest.fail('budget exhausted'))
    class Opener:
        def open(self, req, timeout):
            assert timeout == 0.5
            raise TimeoutError()
    with pytest.raises(TimeoutError):
        sales.read_source(Opener(), sales.REPORT_URL, 100, deadline=100.5)
    with pytest.raises(sales.ReportError, match='source_timeout'):
        sales.read_source(Opener(), sales.REPORT_URL, 100, deadline=99)


def test_partial_response_is_discarded_before_retry(monkeypatch):
    attempts = []
    class Response:
        def __enter__(self): return self
        def __exit__(self, *args): pass
        def read(self, limit):
            if len(attempts) == 1: raise IncompleteRead(b'partial')
            return b'complete'
    class Opener:
        def open(self, *args, **kwargs):
            attempts.append(1)
            return Response()
    monkeypatch.setattr(sales.time, 'sleep', lambda _: None)
    assert sales.read_source(Opener(), sales.REPORT_URL, 100) == b'complete'
    assert len(attempts) == 2


@pytest.mark.parametrize('signature', [b'\x09\x00', b'\x09\x02', b'\x09\x04', b'\x09\x08', b'\xd0\xcf\x11\xe0'])
def test_download_accepts_raw_biff_and_ole_xls(signature, monkeypatch):
    class Response:
        def __enter__(self):return self
        def __exit__(self,*args):pass
        def read(self,*args):return signature+b'workbook'
    class Opener:
        def open(self, req, **kwargs):
            assert b'frmStart=09%2F01%2F2026' in req.data
            assert b'frmEnd=09%2F01%2F2026' in req.data
            return Response()
    monkeypatch.setattr(sales,'read_xls',lambda data:['parsed'])
    assert sales.fetch_report(Opener(),date(2026,9,1))==['parsed']


@pytest.mark.parametrize('path', ['sales', 'sales/export'])
def test_reports_require_login(client, path):
    assert client.get('/api/internal/' + path).status_code == 401


def test_member_can_read_but_only_admin_can_refresh(client, monkeypatch):
    monkeypatch.setattr(sales, 'report_data', lambda *args: {'rows': []})
    calls = []
    def run_batch(**kwargs):
        calls.append(kwargs)
        return {'status': 'complete'}
    monkeypatch.setattr(sales, 'run_batch', run_batch)
    login(client, 'Member')
    assert client.get('/api/internal/sales').status_code == 200
    assert write(client, 'sales/refresh', {}).status_code == 403
    login(client)
    assert write(client, 'sales/refresh', {}).status_code == 200
    assert client.post('/api/internal/sales/refresh', json={}).status_code == 403
    assert calls == [{'force_today': True}]


def test_scheduler_fails_closed_and_blocks_preview(client, monkeypatch):
    monkeypatch.setattr(sales, 'run_batch', lambda: {'status': 'complete'})
    monkeypatch.delenv('CRON_SECRET', raising=False)
    assert client.post('/api/jobs/sales').status_code == 401
    monkeypatch.setenv('CRON_SECRET', 'test-secret')
    assert client.post('/api/jobs/sales', headers={'Authorization':'Bearer wrong'}).status_code == 401
    h = {'Authorization': 'Bearer test-secret'}
    assert client.post('/api/jobs/sales', headers=h).status_code == 200
    monkeypatch.setenv('VERCEL_ENV', 'preview')
    assert client.post('/api/jobs/sales', headers=h).status_code == 403


@pytest.mark.parametrize('query', ['start=invalid', 'start=2026-01-01&end=2026-05-01', 'start=2099-01-01', 'start=2026-09-20&end=2026-09-01'])
def test_invalid_ranges_are_rejected_before_database(client, query):
    login(client)
    assert client.get('/api/internal/sales?' + query).status_code == 400


@pytest.mark.parametrize('force', [False, True])
def test_six_workers_actually_start_together(monkeypatch, force):
    barrier = Barrier(6, timeout=5)
    calls = []
    class Result:
        def __init__(self, value): self.value = value
        def fetchone(self): return self.value
    class Conn:
        def __enter__(self): return self
        def __exit__(self, *args): pass
        def execute(self, sql, *args):
            if 'try_advisory' in sql: return Result({'locked':True})
            return Result({'remaining':0,'due':0,'failed':0})
    def worker(source, cookies, current):
        calls.append(source)
        barrier.wait()
        return {'source':source,'completed':1,'failed':0}
    monkeypatch.setattr(sales, 'connect', Conn)
    seeded = []
    monkeypatch.setattr(sales, 'seed_jobs', lambda *args, **kwargs: seeded.append(kwargs))
    monkeypatch.setattr(sales, 'source_worker', worker)
    assert sales.run_batch(force_today=force)['status'] == 'complete'
    assert seeded == [{'force_today': force}]
    assert set(calls) == set(sales.SOURCES)


@pytest.mark.parametrize('force', [False, True])
def test_busy_lock_does_not_start_downloads(monkeypatch, force):
    class Result:
        def fetchone(self): return {'locked':False}
    class Conn:
        def __enter__(self): return self
        def __exit__(self, *args): pass
        def execute(self, *args): return Result()
    monkeypatch.setattr(sales, 'connect', Conn)
    monkeypatch.setattr(sales, 'source_worker', lambda *args: pytest.fail('must not download'))
    monkeypatch.setattr(sales, 'seed_jobs', lambda *args, **kwargs: pytest.fail('must not requeue'))
    assert sales.run_batch(force_today=force)['status'] == 'busy'


def test_xlsx_keeps_dealer_and_blocks_formula_injection():
    rows = sales.aggregate([row(store='=HYPERLINK("https://example.com")', total_boxes=2, accessory=Decimal('40'),qpay=4)])
    data = {'rows':rows,'totals':rows[0], 'start':'2026-09-01','end':'2026-09-01',
            'coverage':{'complete':6,'expected':6,'retained':0},'updated_at':None}
    book = load_workbook(sales.workbook_bytes(data))
    sheet = book.active
    assert sheet.cell(1,1).value == 'Dealer'
    assert sheet.cell(2,1).value == 'AMQ'
    assert sheet.cell(2,3).data_type == 's'
    assert sheet.cell(2,10).value == 20
    assert sheet.cell(2,13).value == 50


def test_parser_accepts_unassigned_market_and_rejects_duplicate_store(monkeypatch):
    headers=['custno','company','marketid','newact','upgsor','reactact','totact','totaccessory','totpaymentqty','iot','hint']
    values=['S1','Store','',2,1,0,3,60,6,2,1]
    class Sheet:
        nrows=2
        def row_values(self, i):return headers
        def cell_value(self, r, c):return values[c]
    class Book:
        def sheet_by_index(self, i):return Sheet()
    monkeypatch.setattr(sales.xlrd,'open_workbook',lambda **kw:Book())
    result=sales.read_xls(b'dummy')
    assert result[0]['market']=='UNASSIGNED'
    assert result[0]['bts']==1
    Sheet.nrows=3
    with pytest.raises(sales.ReportError,match='duplicate_store'):
        sales.read_xls(b'dummy')
