from datetime import date
from io import BytesIO
import pytest
from openpyxl import load_workbook
from api import index, store_hours as sh, store_hours_sheet as sheet
from api.sheet_types import COMPLIANCE_COLUMNS, CHECK_FIELDS, CHECK_COLORS, sheet_table
from api.portal import PortalError
from test_portal import store, client, login, write
from test_store_hours import sample_rows


def test_compliance_page_and_permission(client, store):
    response=client.get('/compliance')
    assert response.status_code==200 and b'data-sheet="compliance"' in response.data
    assert b'Store Compliance' in response.data
    assert client.get('/api/internal/compliance/sheet').status_code==401
    store.db('portal_users','PATCH',{'page_access':['spreadsheet']},id='eq.'+store.member['id'])
    login(client,'Member')
    for path in ['compliance/sheet','compliance/history','compliance/sheet/export']:
        assert client.get('/api/internal/'+path).status_code==403
    for path in ['compliance/review','compliance/sheet/presence','compliance/refresh','compliance/rows']:
        assert write(client,path,{}).status_code==403


def test_compliance_validation_and_storage_scope():
    with index.app.test_request_context('/api/internal/compliance/review'):
        assert sheet_table('employee_reviews')=='public.compliance_employee_reviews'
        for value in ['Yes','No','#N/A',None,'']:
            assert sh.validate_cell_overrides({'shirt':value})=={'shirt':value or None}
        for value in ['Maybe','true',1,[],{}]:
            with pytest.raises(PortalError): sh.validate_cell_overrides({'shirt':value})
        with pytest.raises(PortalError): sh.validate_cell_overrides({'camera_in':'Camera Offline'})
        sheet.presence_column('shirt')
        with pytest.raises(PortalError): sheet.presence_column('camera_in')
    assert sheet_table('employee_reviews')=='public.store_hours_employee_reviews'
    with pytest.raises(PortalError): sh.validate_cell_overrides({'shirt':'Yes'})


def test_compliance_export_colors_and_text():
    with index.app.test_request_context('/api/internal/compliance/sheet/export'):
        row=sample_rows()[0]
        row.update({k:None for k in CHECK_FIELDS})
        row.update(store_cleaning='No',name_tag='Yes',shirt='#N/A')
        row['cell_formats']={'shirt':{'background':'#FFFFFF','color':'#000000'}}
        data=dict(rows=[row],columns=COMPLIANCE_COLUMNS,start=row['date'],end=row['date'],updated_at=None,coverage={'complete':1,'expected':1})
        book=load_workbook(sh.excel_report(data))
        ws=book.active
        assert ws.title=='Compliance'
        assert [c.value for c in ws[2]]==[label for _,label in COMPLIANCE_COLUMNS]
        for col,value in [(5,'No'),(6,'Yes'),(7,'#N/A')]:
            cell=ws.cell(3,col)
            assert cell.value==value and cell.data_type=='s'
            assert cell.font.color.rgb.endswith(CHECK_COLORS[value][0])
            assert cell.fill.fgColor.rgb.endswith(CHECK_COLORS[value][1])


def test_camera_offline_is_pending_without_chargeback():
    row=sample_rows()[0]
    row.update(camera_in='Camera Offline',camera_out='Camera Offline')
    result=sh.evaluate(row)
    assert result['chargeback'] is None
    assert result['opening_status']==result['closing_status']=='Pending camera review'
    assert sh.display_value('camera_in','Camera Offline')=='Camera Offline'
    assert sh.validate_cell_overrides({'camera_in':'Camera Offline'})=={'camera_in':'Camera Offline'}
    with pytest.raises(PortalError): sh.validate_cell_overrides({'camera_in':'Broken'})
