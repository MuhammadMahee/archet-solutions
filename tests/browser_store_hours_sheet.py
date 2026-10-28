"""Standalone spreadsheet: shared columns, formats, filters, export, and keyboard ranges."""
from pathlib import Path
from openpyxl import load_workbook
from playwright.sync_api import expect
from browser_store_hours import main

ROOT=Path(__file__).resolve().parents[1]


def check(page):
    def cell(address): return page.locator(f'#sh-table [data-address="{address}"]')
    def go(address):
        page.locator('#sh-grid-address').fill(address)
        page.locator('#sh-grid-address').press('Enter')
    def edit(address,value):
        go(address)
        page.keyboard.press('F2')
        page.get_by_label('Edit '+address,exact=True).fill(value)
        page.keyboard.press('Enter')
    def save():
        page.locator('#sheet-save').evaluate('(button)=>button.click()')
        expect(page.locator('#sheet-save-state')).to_have_text('All changes saved')

    assert page.locator('.sidebar').count()==0
    expect(page.locator('#sheet-app')).to_be_visible()
    page.locator('#sheet-add-column').click()
    page.get_by_label('Column name',exact=True).fill('Review Notes')
    page.get_by_role('button',name='Save column',exact=True).click()
    expect(page.locator('thead')).to_contain_text('Review Notes')
    edit('P1','Needs follow-up')
    edit('P2','Checked')
    go('P1')
    page.keyboard.press('Control+b')
    expect(cell('P1')).to_have_css('font-weight','700')
    page.keyboard.press('Control+z')
    expect(cell('P1')).to_have_css('font-weight','400')
    page.keyboard.press('Control+y')
    expect(cell('P1')).to_have_css('font-weight','700')
    page.locator('#sheet-fill-color').fill('#fff0aa')
    page.locator('#sheet-fill-color').dispatch_event('change')
    expect(cell('P1')).to_have_css('background-color','rgb(255, 240, 170)')
    go('O1')
    page.locator('#sheet-number-format').select_option('currency')
    expect(cell('O1')).to_have_text('-$10.00')
    save()
    page.reload()
    expect(cell('P1')).to_have_text('Needs follow-up')
    expect(cell('P1')).to_have_css('font-weight','700')
    expect(cell('P1')).to_have_css('background-color','rgb(255, 240, 170)')
    # Saved presentation and all custom columns are exported.
    page.get_by_role('button',name='File',exact=True).click()
    page.get_by_role('menuitem',name='Download Excel',exact=True).click()
    with page.expect_download() as download:
        page.get_by_role('button',name='Download Excel',exact=True).click()
    path=ROOT/'test-results/standalone-sheet.xlsx'
    download.value.save_as(str(path))
    book=load_workbook(path)
    assert book.active['P2'].value=='Review Notes'
    assert book.active['P3'].value=='Needs follow-up' and book.active['P3'].font.bold
    assert book.active['P3'].fill.fgColor.rgb.endswith('FFF0AA') or book.active['P3'].fill.fgColor.rgb.endswith('fff0aa')
    assert book.active['O3'].number_format=='$#,##0.00;[Red]-$#,##0.00'

    # Default headings expose filters and explicitly prevent deletion.
    page.get_by_role('button',name='Filter Employee column employee',exact=True).click()
    expect(page.get_by_role('button',name='Delete column',exact=True)).to_be_disabled()
    page.get_by_label('Text contains',exact=True).fill('Closing')
    page.get_by_role('button',name='Apply filter',exact=True).click()
    expect(page.locator('#sheet-row-count')).to_contain_text('1 row')
    expect(cell('E1')).to_have_text('CLOSING EMPLOYEE')
    expect(cell('P1')).to_have_text('Checked')
    page.get_by_role('button',name='File',exact=True).click()
    page.get_by_role('menuitem',name='Download Excel',exact=True).click()
    with page.expect_download() as download:
        page.get_by_role('button',name='Download Excel',exact=True).click()
    path=ROOT/'test-results/standalone-filtered.xlsx'
    download.value.save_as(str(path))
    assert load_workbook(path).active.max_row==3
    page.locator('#sheet-filter-toggle').click()
    page.locator('#sheet-clear-filters').click()
    expect(page.locator('#sheet-row-count')).to_contain_text('7 rows')
    page.locator('#sheet-filter-toggle').click()

    # The value bar edits selected cells. Custom names are renameable and removable.
    go('P2')
    page.locator('#sheet-value').fill('Verified in camera')
    page.locator('#sheet-value').press('Enter')
    expect(cell('P2')).to_have_text('Verified in camera')
    save()
    page.get_by_role('button',name='Filter Review Notes column custom_1',exact=True).click()
    page.get_by_role('button',name='Rename column',exact=True).click()
    page.get_by_label('Column name',exact=True).fill('Camera Notes')
    page.get_by_role('button',name='Save column',exact=True).click()
    expect(page.locator('thead')).to_contain_text('Camera Notes')
    # Add enough columns to check AA addressing and ranges beyond the defaults.
    for n in range(12):
        response=page.request.post(page.url.split('/spreadsheet')[0]+'/api/internal/store-hours/sheet/columns',
            headers={'Origin':page.url.split('/spreadsheet')[0],'X-Archet-Request':'1'},
            data={'action':'add','name':'Extra '+str(n),'format':'text','revision':n+2})
        assert response.ok
    page.reload()
    expect(cell('AB1')).to_be_attached()
    edit('AB1','Beyond Z')
    go('AB1:AB3')
    page.keyboard.press('Control+d')
    expect(cell('AB3')).to_have_text('Beyond Z')
    page.keyboard.press('Control+Space')
    expect(page.locator('#sh-table .sh-cell-selected')).to_have_count(7)
    save()
    go('A1')
    page.screenshot(path=str(ROOT/'test-results/standalone-spreadsheet.png'),full_page=True)
    page.get_by_role('button',name='Filter Camera Notes column custom_1',exact=True).click()
    page.get_by_role('button',name='Delete column',exact=True).click()
    page.get_by_role('button',name='Remove column',exact=True).click()
    expect(page.locator('thead')).not_to_contain_text('Camera Notes')
    expect(page.locator('thead')).to_contain_text('Employee')
    page.locator('#sh-history').click()
    expect(page.locator('#sh-history-list')).to_contain_text('Spreadsheet columns')
    page.locator('#sh-history-list summary').first.click()
    expect(page.locator('#sh-history-list')).to_contain_text('Removed (values retained)')
    page.locator('#sh-history-close').click()
    page.set_viewport_size({'width':390,'height':844})
    assert page.evaluate('document.documentElement.scrollWidth <= innerWidth')
    page.screenshot(path=str(ROOT/'test-results/standalone-spreadsheet-mobile.png'),full_page=True)
    print('Standalone sheet browser checks passed: no sidebar, custom columns, default protection, formats, undo, filters, exports, persistence, history, AA ranges, mobile.')


if __name__=='__main__': main(check,spreadsheet=True)
