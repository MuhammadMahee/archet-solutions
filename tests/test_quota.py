"""SPDI formula parity, workbook matching, exports, and access boundaries."""
from copy import deepcopy
from datetime import date
from decimal import Decimal as D
from io import BytesIO
from contextlib import contextmanager
import base64

import pytest
from openpyxl import Workbook, load_workbook
from api import quota
from api.portal import PortalError
from test_portal import store, client, login, write


def roster():
    return [{'dealer': dealer, 'store_id': str(i), 'market': 'RENO', 'store': f'STORE {i}', 'version_id': 'roster-v1'}
            for dealer in ('Connect', 'California') for i in (1, 2, 3)]


def workbook(rows=None, headers=None):
    book = Workbook(); sheet = book.active; sheet.title = 'Goals'
    sheet.append(headers or ['Market','Stores','Voice','BTS','HSI/HINT','Acc','MIM'])
    for row in rows if rows is not None else [['RENO','STORE 1',100,20,10,1000,5]]:
        sheet.append(row)
    buffer = BytesIO(); book.save(buffer); book.close()
    return buffer.getvalue()


def fixture_goals():
    return [{**s, 'voice_goal': D(100), 'bts_goal': D(20), 'hsi_goal': D(10),
             'accessory_goal': D(1000), 'mim_goal': D(5)} for s in roster() if s['dealer']=='Connect']


def fixture_actuals():
    return [{**s, 'new_activation': 80, 'reactivation':10, 'upgrade':999,
             'bts':10, 'hsi':5, 'accessory':D(500), 'incomplete':False, 'stale':False}
            for s in roster() if s['dealer']=='Connect']


def test_legacy_headers_duplicate_sum_and_dealer_isolation():
    data = workbook([[' reno ','store 1',100,20,10,'$1,000.25',None],['RENO','STORE 1',1,2,3,4,5]])
    rows = quota.parse_workbook(data, roster(), 'California')
    assert len(rows) == 1 and rows[0]['dealer'] == 'California'
    assert rows[0]['voice_goal'] == 101 and rows[0]['accessory_goal'] == D('1004.25')
    assert rows[0]['mim_goal'] == 5


def test_template_id_survives_renamed_store_and_legacy_mim_optional():
    data = workbook([['old market','old name',1,2,3,4,'1']], ['Market','Stores','Voice','BTS','HSI','Accessories','Store ID'])
    rows = quota.parse_workbook(data, roster(), 'Connect')
    assert rows[0]['store']=='STORE 1' and rows[0]['market']=='RENO' and rows[0]['mim_goal']==0


@pytest.mark.parametrize('value', ['-1','NaN','Infinity','(25)',True,'=1+2','9999999999999'])
def test_invalid_goals_do_not_silently_become_zero(value):
    with pytest.raises(PortalError):
        quota.parse_workbook(workbook([['RENO','STORE 1',value,0,0,0,0]]),roster(),'Connect')


def test_reject_unmatched_and_ambiguous_names_and_empty_workbooks():
    for data, stores in [(workbook([['RENO','UNKNOWN',1,2,3,4,5]]),roster()),
                         (workbook(),roster()+[{**roster()[0],'store_id':'other'}]),
                         (workbook([]),roster()), (b'not a workbook',roster())]:
        with pytest.raises(PortalError):
            quota.parse_workbook(data, stores, 'Connect')


@pytest.mark.parametrize('month', ['2026-13','2026-9','0001-01',None,[], '2026-09-01'])
def test_bad_months(month):
    with pytest.raises(PortalError): quota.month_value(month)


def test_reference_voice_and_summary_formulas_and_weighted_totals():
    goals, actuals = fixture_goals(), fixture_actuals()
    goals[1]['voice_goal']=D(200)
    tables, elapsed, days = quota.build_tables(goals,actuals,date(2026,9,1),date(2026,9,15))
    assert elapsed==15 and days==30
    voice = tables[0]
    assert voice['rows'][0]['actual']==75  # 80 + 10 - 10 - 5, never includes 999 upgrades
    assert voice['total']['growth']==D(225)/400
    row = next(r for r in tables[3]['rows'] if r['store_id']=='1')
    assert row['quota']==135 and row['actual']==90
    assert row['growth']==D(90)/135
    assert row['acc_remain']==500 and row['per_day']==D(500)/15
    assert row['trend']==1000 and row['achieved']==1
    assert row['overall']==(D(90)/135+1)/2
    assert [r['rank'] for r in tables[3]['rows']]==[1,1,3]
    assert tables[3]['total']['growth']==D(270)/505


def test_rank_resets_per_dealer_and_market():
    goals=fixture_goals(); actuals=fixture_actuals()
    goals.extend([{**goals[0],'dealer':'California'},{**goals[0],'market':'DALLAS','store_id':'9'}])
    actuals.extend([{**actuals[0],'dealer':'California'},{**actuals[0],'store_id':'9'}])
    tables,_,_=quota.build_tables(goals,actuals,date(2026,9,1),date(2026,9,15))
    assert all(r['rank']==1 for r in tables[3]['rows'])


def test_zero_goals_negative_voice_overachievement_and_month_boundaries():
    goals=fixture_goals()[:1]; actuals=fixture_actuals()[:1]
    for key in quota.GOALS: goals[0][key]=D(0)
    actuals[0].update(new_activation=0,reactivation=0)
    tables,_,_=quota.build_tables(goals,actuals,date(2026,9,1),date(2026,10,1))
    assert tables[0]['rows'][0]['actual']==-15
    assert tables[0]['rows'][0]['growth']==0
    assert tables[3]['rows'][0]['per_day']==0
    assert tables[3]['rows'][0]['trend']==500
    future,elapsed,_=quota.build_tables(goals,[],date(2026,10,1),date(2026,9,30))
    assert elapsed==0 and future[0]['rows'][0]['actual']==0
    assert not future[0]['rows'][0]['incomplete']
    goals[0]['voice_goal']=D(5);actuals[0].update(new_activation=35)
    tables,_,_=quota.build_tables(goals,actuals,date(2026,9,1),date(2026,9,15))
    assert tables[0]['rows'][0]['growth']==4 and tables[0]['rows'][0]['remaining']==-15


def test_missing_sales_are_blank_and_partial_totals_are_flagged():
    goals=fixture_goals(); actuals=fixture_actuals()[:1]
    actuals[0]['incomplete']=True
    tables,_,_=quota.build_tables(goals,actuals,date(2026,9,1),date(2026,9,15))
    assert sum(r['actual'] is None for r in tables[0]['rows'])==2
    assert tables[0]['total']['actual']==75 and tables[0]['total']['incomplete']
    assert tables[3]['rows'][-1]['rank'] is None
    empty,_,_=quota.build_tables(goals,[],date(2026,9,1),date(2026,9,15))
    assert empty[0]['total']['actual'] is None and empty[3]['total']['trend'] is None


def test_export_four_tables_number_formats_colors_and_formula_injection():
    goals=fixture_goals();goals[0]['store']='=DANGEROUS()'
    tables,elapsed,days=quota.build_tables(goals,fixture_actuals(),date(2026,9,1),date(2026,9,15))
    data={'tables':tables,'month':'2026-09','today':'2026-09-15','incomplete':0,'elapsed':elapsed,'days':days}
    book=quota.report_workbook(data); buffer=BytesIO();book.save(buffer);book.close()
    book=load_workbook(BytesIO(buffer.getvalue()))
    assert len(book.worksheets)==4
    sheet=book['Voice Goals'];assert sheet['D4'].value==100
    assert sheet['C4'].value=='=DANGEROUS()' and sheet['C4'].data_type=='s'
    assert sheet['G4'].number_format=='0.00%;(0.00%)'
    assert sheet['G4'].fill.fgColor.rgb.endswith('BEF264')
    book.close()


def test_routes_require_login_admin_and_same_origin(client):
    for path in ('quota','quota/excel','quota/template'):
        assert client.get('/api/internal/'+path).status_code==401
    assert write(client,'quota/upload',{}).status_code==401
    login(client,'Member')
    assert write(client,'quota/upload',{}).status_code==403
    assert client.get('/api/internal/quota/template?dealer=Connect').status_code==403
    assert client.post('/api/internal/quota/upload',json={}).status_code==403
    assert client.get('/api/internal/quota',base_url='https://archetsolutions.com').status_code==404


class UploadDB:
    def __init__(self): self.saved={}; self.goals={}; self.args=None
    def __enter__(self): return self
    def __exit__(self,*args): pass
    @contextmanager
    def transaction(self):
        before=deepcopy((self.saved,self.goals))
        try: yield
        except Exception:
            self.saved,self.goals=before;raise
    def execute(self,sql,args=None):
        if sql.startswith('insert into public.quota_uploads'): self.saved[args[:2]]=args[2:]
        elif sql.startswith('delete from public.quota_goals'): self.goals[args]=[]
    def cursor(self): return self
    def executemany(self,sql,rows):
        for row in rows: self.goals[row[:2]].append(row)


def test_upload_preview_then_apply_is_scoped_and_rechecks_roster(client,monkeypatch):
    db=UploadDB();monkeypatch.setattr(quota.sales,'connect',lambda:db);monkeypatch.setattr(quota,'active_roster',lambda _:roster())
    login(client)
    payload={'dealer':'Connect','month':'2026-09','filename':'Goals.xlsx','content':base64.b64encode(workbook()).decode()}
    result=write(client,'quota/upload',payload)
    assert result.status_code==200 and result.json['count']==1 and not db.saved
    assert write(client,'quota/upload',{**payload,'apply':True,'roster_version':'wrong'}).status_code==409
    assert not db.saved
    for dealer in ('Connect','California'):
        assert write(client,'quota/upload',{**payload,'dealer':dealer,'apply':True,'roster_version':'roster-v1'}).status_code==200
    assert len(db.saved)==2 and len(db.goals)==2
    assert write(client,'quota/upload',{**payload,'content':'invalid','apply':True,'roster_version':'roster-v1'}).status_code==400
    assert len(db.saved)==2
    # The upload route allows workbook-sized JSON, but ordinary writes retain their small limit.
    assert write(client,'quota/upload',{**payload,'content':'x'*40000}).status_code==400
    assert write(client,'password',{'padding':'x'*40000}).status_code==413


def test_arbf_accepts_no_acc_column_and_ignores_legacy_acc_values():
    stores=[{**roster()[0], 'dealer':'ARBF'}]
    data=workbook([['RENO','STORE 1',100,20,10]], ['Market','Stores','Voice','BTS','HSI/HINT'])
    assert quota.parse_workbook(data,stores,'ARBF')[0]['accessory_goal']==0
    with pytest.raises(PortalError, match='Acc'):
        quota.parse_workbook(data,roster(),'Connect')
    legacy=workbook([['RENO','STORE 1',100,20,10,'=ignored()',5]])
    assert quota.parse_workbook(legacy,stores,'ARBF')[0]['accessory_goal']==0


def test_arbf_existing_goals_do_not_affect_scores_or_exports():
    goals=[{**g,'dealer':'ARBF'} for g in fixture_goals()]
    actuals=[{**r,'dealer':'ARBF'} for r in fixture_actuals()]
    actuals[1]['accessory']=D(999999)
    tables,elapsed,days=quota.build_tables(goals,actuals,date(2026,9,1),date(2026,9,15))
    summary=tables[-1]
    for row in [*summary['rows'],summary['total']]:
        assert row['overall']==row['growth']==D(90)/135
        assert all(row[key] is None for key in quota.ACCESSORY_TARGET_FIELDS)
    assert [r['rank'] for r in summary['rows']]==[1,1,1]
    assert all(c['key'] not in quota.ACCESSORY_TARGET_FIELDS for c in summary['columns'])
    assert summary['total']['acc_actual']==1000999
    book=quota.report_workbook({'tables':tables,'month':'2026-09','today':'2026-09-15','incomplete':0,'elapsed':elapsed,'days':days})
    headers=[c.value for c in book['Achievement Summary'][3]]
    assert 'Acc Goal' not in headers and 'Achieved' not in headers and 'Acc Actual' in headers
    book.close()


def test_mixed_totals_exclude_arbf_from_accessory_targets_only():
    goals=[fixture_goals()[0],{**fixture_goals()[0],'dealer':'ARBF'}]
    actuals=[fixture_actuals()[0],{**fixture_actuals()[0],'dealer':'ARBF','accessory':D(999999)}]
    tables,_,_=quota.build_tables(goals,actuals,date(2026,9,1),date(2026,9,15))
    summary=tables[-1];total=summary['total']
    assert total['acc_goal']==1000 and total['acc_remain']==500
    assert total['achieved']==1 and total['per_day']==D(500)/15
    assert total['acc_actual']==1000499 and total['trend']==2000998
    assert total['growth']==D(180)/270
    arbf=next(r for r in summary['rows'] if r['dealer']=='ARBF')
    assert arbf['acc_goal'] is None and arbf['overall']==arbf['growth']
    assert any(c['key']=='acc_goal' for c in summary['columns'])


def test_arbf_template_omits_accessory_goal(client,monkeypatch):
    monkeypatch.setattr(quota.sales,'connect',lambda:UploadDB())
    monkeypatch.setattr(quota,'active_roster',lambda _:[{**roster()[0],'dealer':'ARBF'}])
    login(client)
    result=client.get('/api/internal/quota/template?dealer=ARBF')
    assert result.status_code==200
    book=load_workbook(BytesIO(result.data))
    assert [c.value for c in book.active[1]]==['Store ID','Market','Stores','Voice','BTS','HSI/HINT','MIM']
    book.close()
