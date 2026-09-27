"""Exercise Quota Update with real auth/routes/formulas and isolated database boundaries."""
from datetime import date
from decimal import Decimal
from io import BytesIO
from pathlib import Path
import sys
from threading import Thread
from unittest.mock import patch

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
from flask import request
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright, expect
from werkzeug.serving import make_server
from api import index, portal, quota
from test_portal import FakeStore
from test_quota import UploadDB, fixture_goals, fixture_actuals, roster, workbook
from browser_sales import fixture_report as sales_fixture


def fixture_report():
    goals=fixture_goals(); actuals=fixture_actuals()
    goals += [{**g,'dealer':'California'} for g in fixture_goals()]
    actuals += [{**g,'dealer':'California'} for g in fixture_actuals()]
    dealer=request.args.get('dealer','');market=request.args.get('market','');store=request.args.get('store','')
    if dealer=='ARBF':
        goals=[{**g,'dealer':'ARBF'} for g in fixture_goals()]
        actuals=[{**r,'dealer':'ARBF'} for r in fixture_actuals()]
    if dealer: goals=[r for r in goals if r['dealer']==dealer]
    markets=sorted({r['market'] for r in goals})
    if market: goals=[r for r in goals if r['market']==market]
    stores=[{'id':r['dealer']+':'+r['store_id'],'name':r['store'],'dealer':r['dealer']} for r in goals]
    if store: goals=[r for r in goals if r['dealer']+':'+r['store_id']==store]
    month=request.args.get('month','2026-09')
    tables,elapsed,days=quota.build_tables(goals,actuals,quota.month_value(month),date(2026,9,15))
    return {'tables':tables,'month':month,'months':['2026-09','2026-08'],'today':'2026-09-15','elapsed':elapsed,
            'days':days,'count':len(goals),'markets':markets,'stores':stores,'dealers':list(quota.sales.COLORS),
            'uploads':[{'dealer':'Connect','filename':'Goals.xlsx','row_count':3,'uploaded_at':'2026-09-01T12:00:00Z'}],
            'incomplete':0,'calling_tree_active':True,'updated_at':'2026-09-15T18:30:00Z'}


def main():
    fake=FakeStore();db=UploadDB();server=make_server('127.0.0.1',0,index.app,threaded=True)
    origin=f'http://127.0.0.1:{server.server_port}'
    with patch.object(portal,'db',fake.db),patch.object(portal,'auth',fake.auth),patch.object(quota,'report_data',fixture_report),patch.object(quota.sales,'connect',lambda:db),patch.object(quota.sales,'report_data',sales_fixture),patch.object(quota,'active_roster',lambda _:roster()+[{**s,'dealer':'ARBF'} for s in roster() if s['dealer']=='Connect']),patch.object(quota.sales,'run_batch',return_value={'status':'complete','remaining':0}) as refresh:
        Thread(target=server.serve_forever,daemon=True).start()
        try:
            with sync_playwright() as p:
                browser=p.chromium.launch()
                context=browser.new_context(viewport={'width':1536,'height':1000},permissions=['clipboard-read','clipboard-write'])
                page=context.new_page();errors=[];page.on('pageerror',lambda e:errors.append(str(e)))
                page.goto(origin+'/internal#quota')
                page.locator('#login-username').fill('Mahee');page.locator('#login-password').fill('TestPass123');page.locator('#login-form button[type=submit]').click()
                expect(page.locator('#quota-count')).to_have_text('6 STORES')
                assert page.locator('.quota-card').count()==4
                expect(page.locator('#quota-card-voice tbody tr').first.locator('td').nth(4)).to_have_text('75')
                page.locator('#quota-dealer-trigger').click();page.locator('#quota-dealer-menu [role=option]').filter(has_text='Connect').click()
                expect(page.locator('#quota-count')).to_have_text('3 STORES')
                page.locator('#quota-sync').click()
                expect(page.locator('#quota-message')).to_have_text('Sources synced. Quota results reloaded.')
                refresh.assert_called_once_with(force_today=True)
                expect(page.locator('#quota-dealer-trigger')).to_contain_text('Connect')
                refresh.return_value={'status':'busy'}
                page.locator('#quota-sync').click();expect(page.locator('#quota-message')).to_contain_text('already running')
                refresh.return_value={'status':'partial','remaining':2}
                page.locator('#quota-sync').click();expect(page.locator('#quota-message')).to_contain_text('still incomplete')
                refresh.side_effect=portal.PortalError('Source sync unavailable.',503)
                page.locator('#quota-sync').click();expect(page.locator('#quota-message')).to_have_text('Source sync unavailable.')
                expect(page.locator('#quota-sync')).to_be_enabled()
                refresh.side_effect=None
                connect_color=page.evaluate("getComputedStyle(document.body).getPropertyValue('--sales-accent')")
                page.locator('#quota-market-trigger').click();page.locator('#quota-market-menu input').fill('reno');page.locator('#quota-market-menu [role=option]').filter(has_text='RENO').click()
                expect(page.locator('#quota-count')).to_have_text('3 STORES')
                # Canvas assertions ensure Dealer is absent and text has no shadows.
                page.evaluate('''() => {window.drawn=[];const original=CanvasRenderingContext2D.prototype.fillText;CanvasRenderingContext2D.prototype.fillText=function(text,...args){window.drawn.push({text,shadow:this.shadowBlur});return original.call(this,text,...args);};}''')
                page.locator('[data-quota-copy=voice]').click()
                expect(page.locator('#quota-message')).to_contain_text('snapshot copied')
                image=page.evaluate('''async()=>{const items=await navigator.clipboard.read();const blob=await items[0].getType('image/png');const img=await createImageBitmap(blob);return {width:img.width,height:img.height};}''')
                assert image['width']==1920 and image['height']<500
                drawn=page.evaluate('window.drawn');assert all(r['text']!='DEALER' for r in drawn)
                assert all(r['shadow']==0 for r in drawn)
                page.locator('[data-quota-copy=summary]').click()
                expect(page.locator('#quota-message')).to_contain_text('GOALS ACHIEVEMENT SUMMARY snapshot copied')
                png=page.evaluate('''async()=>{const items=await navigator.clipboard.read();const blob=await items[0].getType('image/png');return Array.from(new Uint8Array(await blob.arrayBuffer()));}''')
                (ROOT/'test-results/quota-summary-snapshot.png').write_bytes(bytes(png))
                with page.expect_download() as download:
                    page.locator('#quota-excel').click()
                book=load_workbook(BytesIO(Path(download.value.path()).read_bytes()));assert len(book.worksheets)==4;assert book.worksheets[0].max_row==7;book.close()
                page.locator('#quota-dealer-trigger').click();page.locator('#quota-dealer-menu [role=option]').filter(has_text='California').click()
                expect(page.locator('#quota-count')).to_have_text('3 STORES')
                assert page.evaluate("getComputedStyle(document.body).getPropertyValue('--sales-accent')")!=connect_color
                page.locator('nav [data-page=sales]').click();expect(page.locator('#sales-count')).to_have_text('72 STORES')
                assert page.evaluate("getComputedStyle(document.body).getPropertyValue('--sales-accent')").strip()=='#095570'
                page.locator('nav [data-page=quota]').click();expect(page.locator('#quota-count')).to_have_text('3 STORES')
                assert page.evaluate("getComputedStyle(document.body).getPropertyValue('--sales-accent')").strip()=='#a92d49'
                page.locator('#quota-upload summary').click();page.locator('#quota-upload-dealer').select_option('Connect')
                with page.expect_download() as download:
                    page.locator('#quota-template').click()
                book=load_workbook(BytesIO(Path(download.value.path()).read_bytes()));assert book.active.max_row==4;assert book.active['A2'].value=='1';book.close()
                page.locator('#quota-file').set_input_files({'name':'Goals.xlsx','mimeType':'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet','buffer':workbook()})
                page.locator('#quota-preview-button').click();expect(page.locator('#quota-preview')).to_be_visible()
                assert not db.saved
                page.locator('#quota-save').click();expect(page.locator('#quota-message')).to_contain_text('store goals saved')
                assert len(db.saved)==1;expect(page.locator('#quota-dealer-trigger')).to_contain_text('Connect')
                assert page.locator('#quota-preview').is_hidden()
                page.screenshot(path=str(ROOT/'test-results/quota-desktop.png'),full_page=True)
                # Failed reload clears previous results instead of leaving a mismatched report.
                page.route('**/api/internal/quota?*',lambda route:route.fulfill(status=503,json={'message':'Data unavailable.'}))
                page.locator('#quota-reload').click();expect(page.locator('#quota-message')).to_have_text('Data unavailable.')
                assert page.locator('.quota-card').count()==0 and page.locator('#quota-excel').is_disabled()
                page.unroute('**/api/internal/quota?*');page.locator('#quota-reload').click();expect(page.locator('#quota-count')).to_have_text('3 STORES')
                page.set_viewport_size({'width':390,'height':844});page.screenshot(path=str(ROOT/'test-results/quota-mobile.png'),full_page=True)
                assert page.evaluate('document.documentElement.scrollWidth <= innerWidth')
                page.locator('#quota-dealer-trigger').click();page.locator('#quota-dealer-menu [role=option]').filter(has_text='ARBF').click()
                expect(page.locator('#quota-card-summary th').filter(has_text='Acc Goal')).to_have_count(0)
                expect(page.locator('#quota-card-summary th').filter(has_text='Acc Actual')).to_have_count(1)
                expect(page.locator('#quota-card-summary tbody tr').first.locator('td').last).to_have_text('66.67%')
                page.locator('#quota-upload-dealer').select_option('ARBF')
                no_acc=workbook([['RENO','STORE 1',100,20,10,5]],['Market','Stores','Voice','BTS','HSI/HINT','MIM'])
                page.locator('#quota-file').set_input_files({'name':'ARBF.xlsx','mimeType':'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet','buffer':no_acc})
                page.locator('#quota-preview-button').click();expect(page.locator('#quota-preview')).to_be_visible()
                expect(page.locator('#quota-preview-table th').filter(has_text='Acc')).to_have_count(0)
                page.locator('#quota-save').click();expect(page.locator('#quota-message')).to_contain_text('store goals saved for ARBF')
                assert len(db.saved)==2
                page.evaluate('window.drawn=[]');page.locator('[data-quota-copy=summary]').click()
                expect(page.locator('#quota-message')).to_contain_text('snapshot copied')
                assert 'ACC GOAL' not in [r['text'] for r in page.evaluate('window.drawn')]
                page.locator('#sidebar-toggle').click();page.locator('#logout').click();expect(page.locator('#login')).to_be_visible()
                assert page.locator('.quota-card').count()==0
                page.locator('#login-username').fill('Member');page.locator('#login-password').fill('TestPass123');page.locator('#login-form button[type=submit]').click()
                expect(page.locator('#quota-count')).to_have_text('6 STORES');expect(page.locator('#quota-upload')).to_be_hidden();expect(page.locator('#quota-sync')).to_be_hidden()
                assert not errors,errors
                browser.close()
            print('Quota desktop/mobile, filters, upload, permissions, PNG and Excel passed.')
        finally: server.shutdown()


if __name__=='__main__': main()
