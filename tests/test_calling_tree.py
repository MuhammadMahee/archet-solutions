import base64
from datetime import date
from io import BytesIO
from zipfile import ZipFile

import pytest
from openpyxl import Workbook
from api import calling_tree, sales
from api.portal import PortalError
from test_portal import store,client,login,write


def workbook(rows=None,headers=None):
    book=Workbook();sheet=book.active
    sheet.append(headers or ['Dealer','Store ID','Market','Store Name','DM'])
    for row in rows or [['CONNECT','NTG-REN01','Reno','1940 E WILLIAM','Manager']]:
        sheet.append(row)
    out=BytesIO();book.save(out);return out.getvalue()


def test_dealer_aliases_preserve_ids_and_metadata():
    _,rows=calling_tree.parse_workbook(workbook([
        ['DF Wireless','N1301A','Chicago','Store 1','Person'],
        ['Arm Wireless','317-002','Chicago','Store 2','Person'],
        ['SPDI-CA','SPSCA1','Cali','Store 3','Person'],
        ['SUPREME','TX1','Dallas','Store 4','Person'],
        ['ARBF Wireless Metro','C-70851157','Metro','Store 5','Person']]))
    assert [r['dealer'] for r in rows]==['AMQ','ARM','California','SRH','ARBF']
    assert rows[0]['original_dealer']=='DF Wireless'
    assert rows[0]['store_id']=='N1301A'
    assert rows[0]['market']=='CHICAGO'


@pytest.mark.parametrize('rows,message',[
    ([['Connect','S1','Reno','Store'],['CONNECT','S1','Reno','Other']],'duplicate'),
    ([['Unknown','S1','Reno','Store']],'unknown dealer'),
    ([['AMQ','S1','Reno','Store'],['ARM1','S1','Reno','Other']],'duplicate'),
    ([['Connect','','Reno','Store']],'cannot be empty'),
    ([['Connect','S1','Reno','=HYPERLINK("https://example.com")']],'formulas'),
])
def test_bad_workbooks_fail_before_replacing_active_roster(rows,message):
    with pytest.raises(PortalError,match=message):calling_tree.parse_workbook(workbook(rows))


def test_same_store_id_in_different_dealers_is_not_merged():
    _,rows=calling_tree.parse_workbook(workbook([['AMQ','S1','Reno','One'],['ARM','S1','Reno','Two']]))
    assert len(rows)==2
    assert [r['dealer'] for r in rows] == ['AMQ','ARM']


def test_legacy_arm_names_map_to_new_dealers():
    _,rows=calling_tree.parse_workbook(workbook([['ARM1','S1','Reno','One'],['ARM2','S2','Reno','Two']]))
    assert [r['dealer'] for r in rows] == ['AMQ','ARM']


def test_wrong_file_empty_workbook_and_headers_rejected():
    with pytest.raises(PortalError):calling_tree.parse_workbook(b'not excel')
    with pytest.raises(PortalError,match='Required columns'):
        calling_tree.parse_workbook(workbook(headers=['Dealer','Wrong','Other']))


def test_zip_expansion_is_bounded():
    from zipfile import ZIP_DEFLATED
    data=BytesIO()
    with ZipFile(data,'w',ZIP_DEFLATED) as archive:archive.writestr('large.xml',b'0'*(26*1024*1024))
    with pytest.raises(PortalError,match='unpacked'):calling_tree.parse_workbook(data.getvalue())


def test_multiple_sheets_require_selection():
    b=Workbook()
    for s in [b.active,b.create_sheet('Second')]:
        s.append(['Dealer','Store ID','Market','Store Name']);s.append(['Connect','S1','Reno','Store'])
    out=BytesIO();b.save(out)
    with pytest.raises(PortalError,match='Several worksheets'):calling_tree.parse_workbook(out.getvalue())
    title,rows=calling_tree.parse_workbook(out.getvalue(),'Second');assert title=='Second' and len(rows)==1


def test_upload_is_admin_only_and_protected_by_origin(client):
    assert client.get('/api/internal/calling-tree').status_code==401
    login(client,'Member')
    assert write(client,'calling-tree/preview',{}).status_code==403
    assert write(client,'calling-tree/activate',{}).status_code==403
    login(client)
    assert client.post('/api/internal/calling-tree/preview',json={}).status_code==403
    assert write(client,'calling-tree/preview',{'filename':'bad.xlsx','content':'bad base64'}).status_code==400
    assert write(client,'calling-tree/preview',{'filename':'bad.xls','content':''}).status_code==400


def test_normal_request_size_limit_still_applies(client):
    login(client)
    assert write(client,'sales/refresh',{'padding':'a'*33000}).status_code==413


def test_allowlist_adds_zero_activity_stores_but_excludes_non_roster_sales():
    day=date(2026,9,26)
    roster=[{'dealer':'Connect','store_id':'S1','market':'RENO','store':'FIRST'},
            {'dealer':'Connect','store_id':'S2','market':'RENO','store':'SECOND'}]
    raw=[{'dealer':'Connect','store_id':'S1','market':'WRONG','store':'WRONG','stale':False,
          **sales.ratios({k:2 for k in sales.METRICS})},
         {'dealer':'Connect','store_id':'EXCLUDED',**{k:999 for k in sales.METRICS}}]
    catalog={('Connect','S1'):['connect'],('Connect','S2'):['connect']}
    imports=[{'source_id':'connect','report_date':day,'status':'complete'}]
    rows=sales.restrict_to_roster(raw,roster,catalog,imports,day,day)
    assert len(rows)==2 and all(r['store_id']!='EXCLUDED' for r in rows)
    assert rows[0]['store']=='FIRST' and rows[0]['market']=='RENO'
    assert rows[1]['total_boxes']==0 and rows[1]['no_activity']
    assert sum(r['total_boxes'] for r in rows)==2


def test_missing_reports_and_unknown_ids_do_not_claim_zero_sales():
    day=date(2026,9,26)
    roster=[{'dealer':'Connect','store_id':'NEW','market':'Reno','store':'NEW STORE'}]
    row=sales.restrict_to_roster([],roster,{},[],day,day)[0]
    assert row['total_boxes'] is None and row['unmatched'] and row['incomplete']
    row=sales.restrict_to_roster([],roster,{('Connect','NEW'):['connect']},[],day,day)[0]
    assert row['total_boxes'] is None and not row['unmatched']


def test_no_calling_tree_means_no_sales_exposure():
    day=date(2026,9,26)
    assert sales.restrict_to_roster([{'dealer':'Connect','store_id':'S1'}],[],{},[],day,day)==[]
